"""Independent P2 value/physics audit; reads authority, writes only this folder."""
from pathlib import Path
import calendar
import json
import math
import numbers
import sys
import time
import traceback
import xml.etree.ElementTree as ET
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PIL import Image
from scipy.stats import spearmanr

HERE = Path(__file__).resolve().parent
COPY = HERE.parent
AUTH = COPY.parent.parent
REPORT = {'authority':str(AUTH),'reproduction':str(COPY),
          'integrity_method':'path, bytes, mtime; numerical values; no digests',
          'atol':1e-7,'rtol':1e-11,'checks':[],'failures':[]}

def check(name, ok, **details):
    item = dict(check=name, passed=bool(ok), **details)
    REPORT['checks'].append(item)
    if not ok:
        REPORT['failures'].append(item)

def close(a,b):
    return np.allclose(np.asarray(a,dtype=float),np.asarray(b,dtype=float),atol=1e-7,rtol=1e-11,equal_nan=True)

def maximum(a):
    return float(np.max(np.abs(np.asarray(a,dtype=float)),initial=0))

def readj(p):
    return json.loads(p.read_text(encoding='utf-8-sig'))

def compare_tree(a,b):
    if isinstance(a,dict):
        return isinstance(b,dict) and a.keys()==b.keys() and all(compare_tree(a[k],b[k]) for k in a)
    if isinstance(a,list):
        return isinstance(b,list) and len(a)==len(b) and all(compare_tree(x,y) for x,y in zip(a,b))
    if isinstance(a,numbers.Real) and not isinstance(a,bool):
        return isinstance(b,numbers.Real) and close(a,b)
    return a==b

def compare_csvs():
    total_cells=0; maxdiff=0.; files=[]
    expected=sorted((AUTH/'results').rglob('*.csv'))
    for p in expected:
        rel=p.relative_to(AUTH); q=COPY/rel
        if not q.is_file():
            check(str(rel),False,reason='missing reproduced CSV'); continue
        a=pd.read_csv(p); b=pd.read_csv(q)
        ok=a.shape==b.shape and list(a.columns)==list(b.columns)
        diff=0.
        if ok:
            for c in a:
                if pd.api.types.is_numeric_dtype(a[c]) and pd.api.types.is_numeric_dtype(b[c]):
                    av=a[c].to_numpy(dtype=float); bv=b[c].to_numpy(dtype=float)
                    ok=bool(ok and close(av,bv))
                    finite=np.isfinite(av)&np.isfinite(bv)
                    diff=max(diff,maximum(av[finite]-bv[finite]))
                    total_cells+=len(a)
                else:
                    ok=bool(ok and a[c].fillna('<NA>').equals(b[c].fillna('<NA>')))
        maxdiff=max(maxdiff,diff); files.append(str(rel))
        check('CSV '+str(rel),ok,rows=len(a),columns=len(a.columns),max_abs_difference=diff)
    REPORT['csv_summary']=dict(files=len(files),numeric_cells=total_cells,max_abs_difference=maxdiff)
    a=readj(AUTH/'results/核心指标.json'); b=readj(COPY/'results/核心指标.json')
    for v in (a,b): v.pop('runtime_seconds',None)
    check('core metrics JSON',compare_tree(a,b))
    REPORT['reproduced_key_metrics']=b

def check_main_physics():
    actual_load=pd.read_excel(COPY/'data/附件/附件2.xlsx',sheet_name='小区负载',index_col=0).to_numpy(float)[31:].ravel()
    actual_pv=pd.read_excel(COPY/'data/附件/附件2.xlsx',sheet_name='光伏发电实际功率',index_col=0).to_numpy(float)[31:].ravel()
    variable=pd.read_excel(COPY/'data/附件/附件4.xlsx',index_col=0).to_numpy(float)[31:].ravel()
    typical=pd.read_excel(COPY/'data/附件/附件1.xlsx')
    q1=pd.read_csv(COPY/'results/问题1_逐时结果.csv')
    check('Q1 source inputs',close(q1.load_kw,typical.iloc[:,2]) and close(q1.pv_kw,typical.iloc[:,3]) and close(q1.price_yuan_per_kwh,typical.iloc[:,1]))
    bal=q1.plan_purchase_kwh+q1.pv_used_kwh+q1.discharge_kwh-q1.load_kw/6-q1.charge_kwh
    soc=q1.soc_end_kwh-q1.soc_start_kwh-.9*q1.charge_kwh+q1.discharge_kwh/.9
    check('Q1 independent physical and cost identities',maximum(bal)<1e-7 and maximum(soc)<1e-7 and close(q1.cost_yuan,q1.price_yuan_per_kwh*q1.plan_purchase_kwh) and close(q1.pv_kw/6,q1.pv_used_kwh+q1.pv_curtailed_kwh),balance_max=maximum(bal),soc_residual_max=maximum(soc))
    check('Q1 cyclic/bounds/power/exclusivity',close(q1.soc_start_kwh.iloc[0],6000) and close(q1.soc_end_kwh.iloc[-1],6000) and close(q1.soc_start_kwh.iloc[1:],q1.soc_end_kwh.iloc[:-1]) and q1.soc_end_kwh.between(1200-1e-7,10800+1e-7).all() and q1[['charge_kwh','discharge_kwh']].max().max()<=5000/6+1e-7 and np.minimum(q1.charge_kwh,q1.discharge_kwh).max()<1e-7)
    REPORT['main_physics']=[]
    for name in ('q2','q3','q4_2','q4_3'):
        f=pd.read_csv(COPY/f'results/{name}_逐时结果.csv'); d=pd.read_csv(COPY/f'results/{name}_每日汇总.csv')
        adjusted=name in ('q3','q4_3')
        check(name+' source mapping',close(f.actual_load_kw,actual_load) and close(f.actual_pv_kw,actual_pv) and close(f.price_yuan_per_kwh,variable if name.startswith('q4') else np.tile(typical.iloc[:,1],334)))
        p=f.initial_plan_kwh.to_numpy(); q=f.final_contract_kwh.to_numpy(); c=f.price_yuan_per_kwh.to_numpy()
        bal=q-f.unused_contract_kwh+f.actual_pv_kw/6-f.pv_curtailed_kwh+f.discharge_kwh+f.emergency_kwh-f.actual_load_kw/6-f.charge_kwh
        soc=f.soc_end_kwh-f.soc_start_kwh-.9*f.charge_kwh+f.discharge_kwh/.9
        bill=c*(np.minimum(p,q)+1.5*np.maximum(q-p,0)+.5*np.maximum(p-q,0)) if adjusted else c*p
        alt=c*(p+1.5*np.maximum(q-p,0)+.5*np.maximum(p-q,0))+5*c*f.emergency_kwh if adjusted else bill+5*c*f.emergency_kwh
        check(name+' balance and SOC identities',maximum(bal)<1e-7 and maximum(soc)<1e-7,balance_max=maximum(bal),soc_residual_max=maximum(soc))
        check(name+' both settlement formulas',close(bill,f.settlement_cost_yuan) and close(5*c*f.emergency_kwh,f.emergency_cost_yuan) and close(bill+5*c*f.emergency_kwh,f.cash_cost_yuan) and close(np.asarray(alt).reshape(334,144).sum(axis=1),d.alternative_cash_cost_yuan))
        energy=['initial_plan_kwh','final_contract_kwh','charge_kwh','discharge_kwh','emergency_kwh','unused_contract_kwh','pv_curtailed_kwh']
        check(name+' nonnegative limits exclusivity',f[energy].min().min()>=-1e-7 and f[['charge_kwh','discharge_kwh']].max().max()<=5000/6+1e-7 and np.minimum(f.charge_kwh,f.discharge_kwh).max()<1e-7 and (f.unused_contract_kwh<=q+1e-7).all() and (f.pv_curtailed_kwh<=f.actual_pv_kw/6+1e-7).all())
        check(name+' SOC continuous all slots including midnight',close(f.soc_start_kwh.iloc[1:],f.soc_end_kwh.iloc[:-1]) and f[['soc_start_kwh','soc_end_kwh']].min().min()>=1200-1e-7 and f[['soc_start_kwh','soc_end_kwh']].max().max()<=10800+1e-7)
        dates=pd.date_range('2025-02-01','2025-12-31').strftime('%Y-%m-%d').tolist()
        check(name+' complete dates and slots',f.date.tolist()==[x for x in dates for _ in range(144)] and f.slot.tolist()==list(range(1,145))*334 and d.date.tolist()==dates)
        for dc,fc in [('plan_kwh','initial_plan_kwh'),('final_contract_kwh','final_contract_kwh'),('cash_cost_yuan','cash_cost_yuan'),('emergency_kwh','emergency_kwh'),('charge_kwh','charge_kwh'),('discharge_kwh','discharge_kwh')]:
            check(name+' daily sum '+dc,close(d[dc],f[fc].to_numpy().reshape(334,144).sum(axis=1)))
        REPORT['main_physics'].append(dict(variant=name,slots=len(f),balance_max=maximum(bal),soc_residual_max=maximum(soc),soc_min=float(f.soc_end_kwh.min()),soc_max=float(f.soc_end_kwh.max()),max_power_kw=float(f[['charge_kwh','discharge_kwh']].max().max()*6),cash=float(f.cash_cost_yuan.sum()),alternate_cash=float(np.sum(alt))))

def check_experiments():
    names=[f'{p}_{b}' for p in ('q2','q42') for b in ('safe','point','typical','no_storage')]
    names += [f'{p}_{b}' for p in ('q3','q43') for b in ('0h','12h','6h','2h')]
    suffixes=['alpha_'+a for a in ('0.70','0.75','0.85','0.90','0.95')]+['clusters_'+s for s in ('5','12','16')]+['history_30','history_90','capacity_0.8','capacity_1.2','power_0.8','power_1.2','roundtrip_0.9']+['emergency_'+s for s in ('3','4','6','8')]
    names += [p+'_'+s for p in ('q2','q43') for s in suffixes]
    table=pd.read_csv(COPY/'results/experiments/实验总表.csv')
    check('54 prescribed experiments exactly',len(names)==54 and sorted(table.name)==sorted(names))
    capacities=[]
    for name in names:
        a=readj(AUTH/f'results/experiments/{name}.json'); b=readj(COPY/f'results/experiments/{name}.json')
        for v in (a,b):
            v.pop('runtime_seconds',None); v.pop('source_files',None)
        check(name+' numerical JSON reproduction',compare_tree(a,b))
        f=pd.read_csv(COPY/f'results/experiments/{name}_daily.csv'); ev=f[f.date>='2025-02-01']; bat=b['battery']; cap=bat['capacity_kwh']; eta=bat['eta_c']
        check(name+' 365d initial and cross-day continuity',f.date.tolist()==pd.date_range('2025-01-01','2025-12-31').strftime('%Y-%m-%d').tolist() and close(f.soc_start_kwh.iloc[0],cap*.5) and close(b['initial_soc_kwh'],cap*.5) and close(f.soc_start_kwh.iloc[1:],f.soc_end_kwh.iloc[:-1]))
        dsoc=f.soc_end_kwh-f.soc_start_kwh-eta*f.charge_kwh+f.discharge_kwh/bat['eta_d']
        check(name+' aggregate physical bounds and energy',maximum(dsoc)<1e-7 and f.soc_min_kwh.min()>=cap*.1-1e-7 and f.soc_max_kwh.max()<=cap*.9+1e-7 and f.max_balance_residual_kwh.max()<1e-7,daily_soc_residual_max=maximum(dsoc))
        totals=b['totals']; sumkeys=['cash_cost_yuan','alternative_cash_cost_yuan','plan_kwh','final_contract_kwh','emergency_kwh','unused_contract_kwh','pv_curtailed_kwh','charge_kwh','discharge_kwh','emergency_slots','settlement_yuan','emergency_yuan']
        check(name+' independently sum 334d',all(close(ev[k].sum(),totals[k]) for k in sumkeys) and totals['days']==334 and close(totals['soc_end_kwh'],f.soc_end_kwh.iloc[-1]) and close(f.cash_cost_yuan,f.settlement_yuan+f.emergency_yuan))
        if '_capacity_' in name:
            factor=float(name.rsplit('_',1)[1]); capacities.append(dict(name=name,capacity=cap,minimum=.1*cap,maximum=.9*cap,initial=.5*cap))
            check(name+' normalized 10/90/50 capacity only',close(cap,12000*factor) and bat['power_kw']==5000 and eta==.9 and bat['eta_d']==.9 and b['spec']['alpha']==.8)
        if name.endswith('no_storage'):
            check(name+' actual battery disabled',cap==0 and bat['power_kw']==0 and maximum(f.charge_kwh)==0 and maximum(f.discharge_kwh)==0)
    REPORT['capacity_checks']=capacities
    for prefix,baseline in (('q2','q2_safe'),('q43','q43_6h')):
        base=pd.read_csv(COPY/f'results/experiments/{baseline}_daily.csv')
        for penalty in (3,4,6,8):
            frame=pd.read_csv(COPY/f'results/experiments/{prefix}_emergency_{penalty}_daily.csv')
            cols=['plan_kwh','final_contract_kwh','emergency_kwh','charge_kwh','discharge_kwh','soc_start_kwh','soc_end_kwh','settlement_yuan']
            check(f'{prefix} emergency {penalty} same physical policy recost',close(frame[cols],base[cols]) and close(frame.emergency_yuan,base.emergency_yuan*penalty/5) and close(frame.cash_cost_yuan,base.settlement_yuan+base.emergency_yuan*penalty/5))
    freq=pd.read_csv(COPY/'results/experiments/更新频率比较.csv')
    for row in freq.to_dict('records'):
        base=readj(COPY/f"results/experiments/{row['variant']}_6h.json"); cand=readj(COPY/f"results/experiments/{row['variant']}_{row['candidate']}.json")
        saving=base['totals']['cash_cost_yuan']-cand['totals']['cash_cost_yuan']; extra=(len(cand['spec']['events'])-4)*334
        check(str(row['variant'])+' '+str(row['candidate'])+' frequency arithmetic',close(row['saving_yuan'],saving) and row['extra_events']==extra and close(row['alternative_saving_yuan'],base['totals']['alternative_cash_cost_yuan']-cand['totals']['alternative_cash_cost_yuan']) and (extra<=0 or close(row['breakeven_yuan_per_extra_event'],saving/extra)))
    REPORT['two_hour_results']=freq[freq.candidate=='2h'].to_dict('records')

def check_workbooks():
    counts=[]
    for filename in ('result1.xlsx','result2.xlsx','result3.xlsx','result4-2.xlsx','result4-3.xlsx'):
        a=load_workbook(AUTH/'results'/filename,read_only=False,data_only=False)
        b=load_workbook(COPY/'results'/filename,read_only=False,data_only=False)
        template=load_workbook(COPY/'data/附件/附件5'/filename,read_only=False,data_only=False)
        ok=a.sheetnames==b.sheetnames==template.sheetnames; numeric=filled=0; md=0.
        for sheet in a.sheetnames:
            x=a[sheet]; y=b[sheet]; t=template[sheet]
            ok=ok and (x.max_row,x.max_column)==(y.max_row,y.max_column) and str(x.merged_cells)==str(y.merged_cells)==str(t.merged_cells)
            ok=ok and [c.value for c in y[1]]==[c.value for c in t[1]]
            for row in x:
                for cell in row:
                    v=cell.value; other=y.cell(cell.row,cell.column)
                    if v is not None: filled+=1
                    ok=ok and other.data_type not in ('e','f')
                    if isinstance(v,numbers.Real) and not isinstance(v,bool):
                        numeric+=1; ok=ok and isinstance(other.value,numbers.Real) and close(v,other.value)
                        if isinstance(other.value,numbers.Real): md=max(md,abs(v-other.value))
                    else: ok=ok and v==other.value
        check('XLSX all cells and template structure '+filename,ok,filled_cells=filled,numeric_cells=numeric,max_abs_difference=md)
        counts.append(dict(file=filename,filled_cells=filled,numeric_cells=numeric,max_abs_difference=md))
        a.close();b.close();template.close()
    REPORT['workbook_summary']=counts

def check_representatives_and_statistics():
    frames={n:pd.read_csv(COPY/f'results/{n}_逐时结果.csv') for n in ('q2','q3','q4_2','q4_3')}
    table=pd.read_csv(COPY/'results/表1_代表日购电.csv')
    for _,r in table.iterrows():
        f=frames[r.variant]; f=f[f.date==r.date].reset_index(drop=True)
        field='initial_plan_kwh' if r.quantity=='initial_plan' else 'final_contract_kwh'
        ok=all(close(r[f'{h}:00-{h}:10_kwh'],f.loc[h*6,field]) for h in (10,12,14,16,18,20))
        check('table1 '+r.variant+' '+r.date+' '+r.quantity,ok and close(r.day_purchase_kwh,f[field].sum()) and close(r.actual_cash_cost_yuan,f.cash_cost_yuan.sum()))
    table=pd.read_csv(COPY/'results/表2_代表日储能.csv')
    for _,r in table.iterrows():
        f=frames[r.variant]; f=f[f.date==r.date].reset_index(drop=True); h=int(r.period.split(':')[0]); cut=f.iloc[h*6:(h+4)*6]
        check('table2 '+r.variant+' '+r.date+' '+r.period,close([r.charge_kwh,r.discharge_kwh,r.soc_0_kwh,r.soc_24_kwh],[cut.charge_kwh.sum(),cut.discharge_kwh.sum(),f.soc_start_kwh.iloc[0],f.soc_end_kwh.iloc[-1]]))
    table=pd.read_csv(COPY/'results/表3_代表日应急.csv')
    for name,f in frames.items():
        for day in ('2025-03-20','2025-06-21','2025-09-23','2025-12-21'):
            expected=f[(f.date==day)&(f.emergency_kwh>1e-8)]
            actual=table[(table.variant==name)&(table.date==day)]
            ok=(actual.interval.tolist()==expected.natural_interval.tolist() and close(actual.emergency_kwh,expected.emergency_kwh)) if len(expected) else len(actual)==1 and actual.interval.iloc[0]=='无' and actual.emergency_kwh.iloc[0]==0
            check('table3 complete emergency slots '+name+' '+day,ok,slots=len(expected))
    q1=pd.read_csv(COPY/'results/问题1_逐时结果.csv')
    stats=pd.read_csv(COPY/'results/图表统计.csv').set_index(['figure','metric']).value
    check('q1 displayed rho recomputed',close(stats.loc[('process_q1_price_dispatch','spearman_rho')],spearmanr(q1.price_yuan_per_kwh,q1.discharge_kwh-q1.charge_kwh).statistic))
    f=frames['q4_3']
    check('q4 displayed descriptive rho recomputed',close(stats.loc[('process_q4_price_dispatch','spearman_rho')],spearmanr(f.price_yuan_per_kwh,f.discharge_kwh-f.charge_kwh).statistic))
    f=frames['q3']; f=f[f.date=='2025-06-21']
    check('q3 displayed absolute adjustment recomputed',close(stats.loc[('process_q3_contract_updates','absolute_adjustment_mwh')],abs(f.final_contract_kwh-f.initial_plan_kwh).sum()/1000))
    daily0=pd.read_csv(COPY/'results/experiments/q3_0h_daily.csv').query('date >= "2025-02-01"')
    daily6=pd.read_csv(COPY/'results/experiments/q3_6h_daily.csv').query('date >= "2025-02-01"')
    check('q3 paired-day benefit statistic',close(stats.loc[('result_q3_adjustment_benefit','days_cost_reduced_percent')],100*np.mean(daily6.cash_cost_yuan.to_numpy()<daily0.cash_cost_yuan.to_numpy())))
    for key,first,second in (('fixed_price_saving_percent','q2','q3'),('variable_price_saving_percent','q4_2','q4_3')):
        base=frames[first].cash_cost_yuan.sum(); adjusted=frames[second].cash_cost_yuan.sum()
        check('q4 total strategy difference '+key,close(stats.loc[('result_q4_strategy_comparison',key)],100*(base-adjusted)/base))

def check_forecast_diagnostics_independently():
    typical=pd.read_excel(COPY/'data/附件/附件1.xlsx')
    loads=pd.read_excel(COPY/'data/附件/附件2.xlsx',sheet_name='小区负载',index_col=0).to_numpy(float)
    pv=pd.read_excel(COPY/'data/附件/附件2.xlsx',sheet_name='光伏发电实际功率',index_col=0).to_numpy(float)
    metrics=pd.read_csv(COPY/'results/预测误差指标.csv')
    stats=pd.read_csv(COPY/'results/图表统计.csv').set_index(['figure','metric']).value
    def metric_values(truth,pred,threshold):
        finite=np.isfinite(truth)&np.isfinite(pred); truth=truth[finite]; pred=pred[finite]
        e=truth-pred; valid=truth>threshold
        return dict(count=len(e),mae_kw=np.mean(abs(e)),rmse_kw=np.sqrt(np.mean(e**2)),bias_kw=np.mean(e),mape_percent=100*np.mean(abs(e[valid])/truth[valid]),mape_count=int(valid.sum()))
    for label,actual,column,weekday,threshold in (('analog_load',loads,2,True,1),('analog_pv',pv,3,False,100)):
        predictions=[]
        for d in range(365):
            previous=list(range(max(0,d-90),d))
            if previous:
                weights=np.array([math.exp(-(d-k)/60)*(2 if weekday and d%7==k%7 else 1) for k in previous])
                weights/=weights.sum(); shrink=min(1,len(previous)/21)
                predictions.append(shrink*np.sum(actual[previous]*weights[:,None],axis=0)+(1-shrink)*typical.iloc[:,column].to_numpy(float))
            else: predictions.append(typical.iloc[:,column].to_numpy(float))
        pred=np.array(predictions); computed=metric_values(actual[31:],pred[31:],threshold)
        row=metrics[metrics.model==label].iloc[0]
        check('independent raw-data forecast metrics '+label,all(close(row[k],v) for k,v in computed.items()),**{k:float(v) for k,v in computed.items()})
        residual=actual-pred; inside=[]; widths=[]
        for d in range(31,365):
            history=np.sort(residual[max(0,d-60):d],axis=0); n=len(history)
            lo=np.maximum(0,pred[d]+history[math.ceil(.05*n)-1]); hi=np.maximum(0,pred[d]+history[math.ceil(.95*n)-1])
            inside.extend(((actual[d]>=lo)&(actual[d]<=hi)).tolist());widths.extend((hi-lo).tolist())
        check('independent inverse empirical interval '+label,close(row.empirical_90_coverage,np.mean(inside)) and close(row.interval_mean_width_kw,np.mean(widths)))
        if label=='analog_load':
            r2=1-np.sum((actual[31:]-pred[31:])**2)/np.sum((actual[31:]-actual[31:].mean())**2)
            check('independent raw-data plotted forecast MAE RMSE R2',all(close(stats.loc[('process_q2_forecast_validation',k)],v) for k,v in (('mae_kw',computed['mae_kw']),('rmse_kw',computed['rmse_kw']),('r2',r2))))
    official=pd.read_excel(COPY/'data/附件/附件3.xlsx')
    dates=pd.to_datetime(official.iloc[:,0].ffill()); raw=np.empty((365,4,24))
    for i,row in enumerate(official.itertuples(index=False,name=None)):
        day=(dates.iloc[i]-pd.Timestamp('2025-01-01')).days
        hour=row[1].hour if hasattr(row[1],'hour') else int(str(row[1]).split(':')[0])
        raw[day,hour//6]=np.array(row[2:],dtype=float)
    truth=np.full(raw.shape,np.nan);flat=pv.ravel()
    for day in range(365):
        for issue in range(4):
            for lead in range(1,25):
                index=day*144+issue*36+lead*6-1
                if index<len(flat): truth[day,issue,lead-1]=flat[index]
    residual=truth-raw
    for _,row in metrics[metrics.model.str.startswith('official')].iterrows():
        if row.model=='official_pv_raw_lead':
            j=int(row.lead_hour)-1; t=truth[31:,:,j];p=raw[31:,:,j]
        else:
            j=int(row.issue_hour)//6; h=24-int(row.issue_hour) if row.model!='official_pv_raw' else 24
            t=truth[31:,j,:h];p=raw[31:,j,:h].copy()
            if row.model=='official_pv_bias_corrected':
                for d in range(31,365):p[d-31]=np.maximum(0,p[d-31]+np.mean(residual[max(0,d-60):d,j,:h],axis=0))
        computed=metric_values(t,p,100)
        check('independent official hourly target metrics '+row.model+' '+str(row.issue_hour)+' '+str(row.lead_hour),all(close(row[k],v) for k,v in computed.items()),count=int(computed['count']))

def check_figures_and_manifest():
    names=sorted(p.stem for p in (AUTH/'figures').glob('*.png'))
    check('16 logical figures and q1-q4 raw/process/result coverage',len(names)==16 and all(any(n.startswith(kind+'_'+q+'_') for n in names) for q in ('q1','q2','q3','q4') for kind in ('raw','process','result')))
    REPORT['figure_pixels']=[]
    for name in names:
        sizes=[]; diffs=[]
        for rel in (Path('figures')/(name+'.png'),Path('figures/_qa')/(name+'_grayscale.png')):
            with Image.open(AUTH/rel) as a, Image.open(COPY/rel) as b:
                av=np.asarray(a);bv=np.asarray(b); eq=av.shape==bv.shape and np.array_equal(av,bv)
                diff=maximum(av.astype(float)-bv.astype(float)) if av.shape==bv.shape else -1
                sizes.append(list(b.size));diffs.append(diff)
                check('decoded pixels '+str(rel),eq,size=list(b.size),max_abs_pixel_difference=diff,dpi=list(b.info.get('dpi',[])))
        svg=ET.parse(COPY/'figures'/(name+'.svg')).getroot()
        images=[x for x in svg.iter() if x.tag.endswith('}image')]
        check('SVG is vector '+name,len(images)==0)
        REPORT['figure_pixels'].append(dict(name=name,sizes=sizes,max_differences=diffs))
    m=readj(COPY/'results/复现清单.json')
    check('full unique entry completed all steps',m['last_invocation_mode']=='full_reproduction' and len(m['steps'])==16 and all(x['exit_code']==0 for x in m['steps']) and any(x['step']=='main' for x in m['steps']) and any(x['step']=='experiments' for x in m['steps']))
    check('manifest seed parameters and input inventory',m['random_seed']==2026 and m['key_parameters']['alpha']==.8 and m['key_parameters']['experiment_cases']==54 and len(m['input_files'])==10 and bool(m['runtime']['dependencies']))
    for item in m['input_files']+m['source_files']:
        p=Path(item['path']);p=p if p.is_absolute() else COPY/p
        check('manifest metadata '+str(p.relative_to(COPY)),p.exists() and p.stat().st_size==item['bytes'] and p.stat().st_mtime_ns==item['mtime_ns'])
    REPORT['pipeline']=dict(mode=m['last_invocation_mode'],runtime_seconds=m['runtime_seconds'],steps=m['steps'],runtime=m['runtime'])

def check_frozen_metadata():
    original=readj(HERE/'input_snapshot.json'); changes=[]
    for r in original:
        p=AUTH/r['path']; stamp=r['mtime_utc'].rstrip('Z'); sec,_,fraction=stamp.partition('.')
        expected=calendar.timegm(time.strptime(sec,'%Y-%m-%dT%H:%M:%S'))*10**9+int((fraction+'000000000')[:9])
        if not p.exists() or p.stat().st_size!=r['bytes'] or p.stat().st_mtime_ns!=expected:
            changes.append(dict(path=r['path'],expected_bytes=r['bytes'],expected_mtime_ns=expected,current_bytes=p.stat().st_size if p.exists() else None,current_mtime_ns=p.stat().st_mtime_ns if p.exists() else None))
    check('authority snapshot path size mtime unchanged',not changes,files=len(original),changes=changes)
    REPORT['authority_metadata_changes']=changes

def main():
    sections=[compare_csvs,check_main_physics,check_experiments,check_workbooks,check_representatives_and_statistics,check_forecast_diagnostics_independently,check_figures_and_manifest,check_frozen_metadata]
    for fn in sections:
        try: fn()
        except Exception as exc:
            check(fn.__name__,False,error=str(exc),traceback=traceback.format_exc())
    REPORT['status']='PASS' if not REPORT['failures'] else 'FAIL'
    (HERE/'numeric_physical_comparison.json').write_text(json.dumps(REPORT,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:REPORT[k] for k in ('status','csv_summary','workbook_summary','failures') if k in REPORT},ensure_ascii=False,indent=2))
    return 0 if not REPORT['failures'] else 1

if __name__=='__main__':
    raise SystemExit(main())
