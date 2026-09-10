"""Full reproducible pipeline for 2026 CUMCM C problem."""

from __future__ import annotations

import argparse
from copy import copy
import json
import platform
from pathlib import Path
import shutil
import sys
import time
from typing import Any

import numpy as np
import pandas as pd
import scipy
import sklearn
from openpyxl import load_workbook
from evaluate_results import forecast_diagnostics, validate_day, write_representative_tables, raw_data_profile

from microgrid_core import (
    DEFAULT_ALPHA,
    DT_HOURS,
    ISSUES,
    N_SLOTS,
    SEED,
    SOC_INITIAL,
    build_analog_forecasts,
    load_inputs,
    official_pv_residuals,
    run_q1,
    run_q2_day,
    run_q3_day,
    safe_analog_trajectory,
    safe_update_trajectory,
    summarize_day,
)


OUTPUT_START = pd.Timestamp("2025-02-01")
REPRESENTATIVE_DATES = (
    pd.Timestamp("2025-03-20"),
    pd.Timestamp("2025-06-21"),
    pd.Timestamp("2025-09-23"),
    pd.Timestamp("2025-12-21"),
)


def natural_interval_labels() -> list[str]:
    labels = []
    for t in range(N_SLOTS):
        a = t * 10
        b = (t + 1) * 10
        labels.append(f"{a // 60:02d}:{a % 60:02d}-{b // 60:02d}:{b % 60:02d}")
    return labels


def file_metadata(path: Path) -> dict:
    """Record file identity without calculating or comparing any hashes."""
    stat = path.stat()
    return {"path": str(path.resolve()), "size_bytes": stat.st_size,
            "modified_time_ns": stat.st_mtime_ns}


def precompute_safe_trajectories(
    data,
    load_point,
    pv_point,
    load_residual,
    pv_residual,
    issued_residual,
    alpha: float,
):
    q2_cache = []
    q3_cache = []
    for day in range(len(data.dates)):
        q2_cache.append(
            safe_analog_trajectory(
                day,
                load_point[day],
                pv_point[day],
                load_residual,
                pv_residual,
                alpha=alpha,
            )
        )
        updates = []
        for issue_index in range(len(ISSUES)):
            updates.append(
                safe_update_trajectory(
                    data,
                    day,
                    issue_index,
                    load_point[day],
                    load_residual,
                    issued_residual,
                    alpha=alpha,
                )
            )
        q3_cache.append(updates)
        if day % 60 == 0 or day == len(data.dates) - 1:
            print(f"[cache] {day + 1}/{len(data.dates)} days", flush=True)
    return q2_cache, q3_cache


def simulate_variant(
    data,
    *,
    adjusted: bool,
    variable_price: bool,
    load_point,
    pv_point,
    load_residual,
    pv_residual,
    issued_residual,
    q2_cache,
    q3_cache,
    alpha: float,
) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    soc = SOC_INITIAL
    name = ("q3" if adjusted else "q2") + ("_variable" if variable_price else "_fixed")
    for day in range(len(data.dates)):
        if adjusted:
            result = run_q3_day(
                data,
                day,
                soc,
                load_point,
                load_residual,
                issued_residual,
                variable_price=variable_price,
                alpha=alpha,
                safe_updates=q3_cache[day],
            )
        else:
            result = run_q2_day(
                data,
                day,
                soc,
                load_point,
                pv_point,
                load_residual,
                pv_residual,
                variable_price=variable_price,
                alpha=alpha,
                safe_pair=q2_cache[day],
            )
        result["date"] = data.dates[day]
        result["day_index"] = day
        validate_day(result, adjusted, soc)
        soc = float(result["actual"]["soc_kwh"][-1])
        results.append(result)
        if day % 45 == 0 or day == len(data.dates) - 1:
            print(f"[{name}] {day + 1}/{len(data.dates)} days, SOC={soc:.2f}", flush=True)
    return results


def daily_summary_frame(results: list[dict[str, Any]], *, adjusted: bool, output_only: bool = True) -> pd.DataFrame:
    rows = []
    for result in results:
        if output_only and result["date"] < OUTPUT_START:
            continue
        row = {"date": result["date"].date().isoformat(), **summarize_day(result, adjusted=adjusted)}
        rows.append(row)
    return pd.DataFrame(rows)


def interval_frame(data, results: list[dict[str, Any]], *, adjusted: bool, output_only: bool = True) -> pd.DataFrame:
    natural = natural_interval_labels()
    rows = []
    for result in results:
        day = result["day_index"]
        if output_only and result["date"] < OUTPUT_START:
            continue
        price = data.variable_price[day] if result.get("variable_price", False) else data.typical_price
        actual = result["actual"]
        plan = result["initial_plan_kwh"] if adjusted else result["plan"]["grid_kwh"]
        contract = result["final_contract_kwh"] if adjusted else plan
        template_labels = result["template_labels"]
        for t in range(N_SLOTS):
            rows.append(
                {
                    "date": result["date"].date().isoformat(),
                    "slot": t + 1,
                    "natural_interval": natural[t],
                    "official_template_interval": template_labels[t],
                    "price_yuan_per_kwh": price[t],
                    "initial_planning_price_yuan_per_kwh": result["planning_price"][t],
                    "actual_load_kw": data.actual_load_kw[day, t],
                    "actual_pv_kw": data.actual_pv_kw[day, t],
                    "initial_plan_kwh": plan[t],
                    "final_contract_kwh": contract[t],
                    "charge_kwh": actual["charge_kwh"][t],
                    "discharge_kwh": actual["discharge_kwh"][t],
                    "soc_start_kwh": actual["soc_kwh"][t],
                    "soc_end_kwh": actual["soc_kwh"][t + 1],
                    "emergency_kwh": actual["emergency_kwh"][t],
                    "unused_contract_kwh": actual["unused_contract_kwh"][t],
                    "pv_curtailed_kwh": actual["pv_curtailed_kwh"][t],
                    "settlement_cost_yuan": actual["settlement_cost"][t],
                    "emergency_cost_yuan": actual["emergency_cost"][t],
                    "cash_cost_yuan": actual["cash_cost"][t],
                }
            )
    return pd.DataFrame(rows)


def q1_frames(data, q1) -> tuple[pd.DataFrame, pd.DataFrame]:
    natural = natural_interval_labels()
    schedule = pd.DataFrame(
        {
            "slot": np.arange(1, N_SLOTS + 1),
            "natural_interval": natural,
            "input_time_label": [str(x) for x in data.typical_time],
            "price_yuan_per_kwh": data.typical_price,
            "load_kw": data.typical_load_kw,
            "pv_kw": data.typical_pv_kw,
            "plan_purchase_kwh": q1["grid_kwh"],
            "charge_kwh": q1["charge_kwh"],
            "discharge_kwh": q1["discharge_kwh"],
            "soc_start_kwh": q1["soc_kwh"][:-1],
            "soc_end_kwh": q1["soc_kwh"][1:],
            "pv_used_kwh": q1["pv_used_kwh"],
            "pv_curtailed_kwh": data.typical_pv_kw * DT_HOURS - q1["pv_used_kwh"],
            "cost_yuan": data.typical_price * q1["grid_kwh"],
        }
    )
    groups = []
    for g in range(6):
        sl = slice(g * 24, (g + 1) * 24)
        groups.append(
            {
                "period": f"{g * 4}:00-{(g + 1) * 4}:00",
                "charge_kwh": float(q1["charge_kwh"][sl].sum()),
                "discharge_kwh": float(q1["discharge_kwh"][sl].sum()),
            }
        )
    aggregate = pd.DataFrame(groups)
    aggregate["soc_0_kwh"] = [q1["soc_kwh"][0]] + [np.nan] * 5
    aggregate["soc_24_kwh"] = [q1["soc_kwh"][-1]] + [np.nan] * 5
    return schedule, aggregate


def _copy_row_style(ws, source_row: int, target_row: int, max_col: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)


def _clear_rows(ws, start: int, end: int, max_col: int) -> None:
    for row in range(start, end + 1):
        for col in range(1, max_col + 1):
            ws.cell(row, col).value = None


def _write_purchase_sheet(ws, data, results, *, adjusted_sheet: bool, total_mode: str) -> list[str]:
    labels = [ws.cell(1, col).value for col in range(2, 2 + N_SLOTS)]
    out = [r for r in results if r["date"] >= OUTPUT_START]
    if len(out) != 334:
        raise ValueError(f"输出日期应为334天，实际{len(out)}")
    for row, result in enumerate(out, start=2):
        plan = result["initial_plan_kwh"] if "initial_plan_kwh" in result else result["plan"]["grid_kwh"]
        contract = result["final_contract_kwh"] if adjusted_sheet else plan
        ws.cell(row, 1).value = result["date"].to_pydatetime()
        for t, value in enumerate(contract, start=2):
            ws.cell(row, t).value = float(value)
        ws.cell(row, 146).value = float(np.sum(contract))
        if total_mode == "actual_cash":
            total_cost = float(np.sum(result["actual"]["cash_cost"]))
        elif total_mode == "base_plan":
            day = result["day_index"]
            price = data.variable_price[day] if result["variable_price"] else data.typical_price
            total_cost = float(np.dot(price, plan))
        else:
            raise ValueError(total_mode)
        ws.cell(row, 147).value = total_cost
        result["template_labels"] = labels
    return labels


def _write_storage_sheet(ws, results) -> None:
    out = [r for r in results if r["date"] >= OUTPUT_START]
    target_last = 1 + 6 * len(out)
    max_col = 6
    for offset, result in enumerate(out):
        actual = result["actual"]
        for g in range(6):
            row = 2 + offset * 6 + g
            _copy_row_style(ws, 2 + (g % 6), row, max_col)
            sl = slice(g * 24, (g + 1) * 24)
            ws.cell(row, 1).value = result["date"].to_pydatetime() if g == 0 else None
            ws.cell(row, 2).value = f"{g * 4}:00-{(g + 1) * 4}:00"
            ws.cell(row, 3).value = float(np.sum(actual["charge_kwh"][sl]))
            ws.cell(row, 4).value = float(np.sum(actual["discharge_kwh"][sl]))
            ws.cell(row, 5).value = "0:00" if g == 0 else ("24:00" if g == 1 else None)
            ws.cell(row, 6).value = float(actual["soc_kwh"][0]) if g == 0 else (float(actual["soc_kwh"][-1]) if g == 1 else None)
    if ws.max_row > target_last:
        _clear_rows(ws, target_last + 1, ws.max_row, max_col)


def _write_emergency_sheet(ws, results) -> None:
    labels = natural_interval_labels()
    rows = []
    for result in results:
        if result["date"] < OUTPUT_START:
            continue
        emergency = result["actual"]["emergency_kwh"]
        for t in np.flatnonzero(emergency > 1e-8):
            rows.append((result["date"].to_pydatetime(), labels[int(t)], float(emergency[t])))
    if not rows:
        rows = [(OUTPUT_START.to_pydatetime(), "无", 0.0)]
    for i, values in enumerate(rows, start=2):
        _copy_row_style(ws, 2 + ((i - 2) % 3), i, 3)
        for col, value in enumerate(values, start=1):
            ws.cell(i, col).value = value
    if ws.max_row > len(rows) + 1:
        _clear_rows(ws, len(rows) + 2, ws.max_row, 3)


def write_result1(data, q1, template: Path, output: Path) -> None:
    shutil.copy2(template, output)
    wb = load_workbook(output)
    ws = wb["计划购电量"]
    for t, value in enumerate(q1["grid_kwh"], start=2):
        ws.cell(t, 2).value = float(value)
    ws = wb["充放电量"]
    for g in range(6):
        sl = slice(g * 24, (g + 1) * 24)
        ws.cell(g + 2, 2).value = float(np.sum(q1["charge_kwh"][sl]))
        ws.cell(g + 2, 3).value = float(np.sum(q1["discharge_kwh"][sl]))
    ws.cell(2, 5).value = float(q1["soc_kwh"][0])
    ws.cell(3, 5).value = float(q1["soc_kwh"][-1])
    wb.save(output)


def write_annual_workbook(data, results, template: Path, output: Path, *, adjusted: bool) -> None:
    shutil.copy2(template, output)
    wb = load_workbook(output)
    _write_purchase_sheet(wb["计划购电量"], data, results, adjusted_sheet=False, total_mode="base_plan" if adjusted else "actual_cash")
    if adjusted:
        _write_purchase_sheet(wb["调整购电量"], data, results, adjusted_sheet=True, total_mode="actual_cash")
    _write_storage_sheet(wb["充放电量"], results)
    _write_emergency_sheet(wb["紧急购电量"], results)
    wb.save(output)


def forecast_metrics(data, load_point, pv_point, issued_residual) -> pd.DataFrame:
    start = int(np.flatnonzero(data.dates == OUTPUT_START)[0])
    rows = []
    for name, truth, pred in (
        ("analog_load", data.actual_load_kw[start:], load_point[start:]),
        ("analog_pv", data.actual_pv_kw[start:], pv_point[start:]),
    ):
        err = truth - pred
        rows.append({"model": name, "issue_hour": np.nan, "lead_hour": np.nan, "mae_kw": np.mean(np.abs(err)), "rmse_kw": np.sqrt(np.mean(err**2)), "bias_kw": np.mean(err)})
    for issue_index, issue in enumerate(ISSUES):
        err = issued_residual[:364, issue_index, :]
        valid = np.isfinite(err)
        rows.append({"model": "official_pv", "issue_hour": issue, "lead_hour": np.nan, "mae_kw": np.mean(np.abs(err[valid])), "rmse_kw": np.sqrt(np.mean(err[valid] ** 2)), "bias_kw": np.mean(err[valid])})
    for lead in (1, 6, 12, 18, 24):
        err = issued_residual[:, :, lead - 1]
        valid = np.isfinite(err)
        rows.append({"model": "official_pv", "issue_hour": np.nan, "lead_hour": lead, "mae_kw": np.mean(np.abs(err[valid])), "rmse_kw": np.sqrt(np.mean(err[valid] ** 2)), "bias_kw": np.mean(err[valid])})
    return pd.DataFrame(rows)


def attach_variant_metadata(results, *, variable_price: bool, template_labels: list[str] | None = None) -> None:
    for result in results:
        result["variable_price"] = variable_price
        if template_labels is not None:
            result["template_labels"] = template_labels


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", default=".")
    parser.add_argument("--alpha", type=float, default=DEFAULT_ALPHA)
    parser.add_argument("--skip-csv", action="store_true")
    args = parser.parse_args()
    root = Path(args.project_root).resolve()
    results_dir = root / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    started = time.time()
    np.random.seed(SEED)

    data = load_inputs(root)
    load_point, load_residual = build_analog_forecasts(data.actual_load_kw, data.typical_load_kw, weekday_weight=True)
    pv_point, pv_residual = build_analog_forecasts(data.actual_pv_kw, data.typical_pv_kw, weekday_weight=False)
    issued_residual = official_pv_residuals(data)
    q2_cache, q3_cache = precompute_safe_trajectories(
        data, load_point, pv_point, load_residual, pv_residual, issued_residual, args.alpha
    )

    q1 = run_q1(data)
    q2 = simulate_variant(data, adjusted=False, variable_price=False, load_point=load_point, pv_point=pv_point, load_residual=load_residual, pv_residual=pv_residual, issued_residual=issued_residual, q2_cache=q2_cache, q3_cache=q3_cache, alpha=args.alpha)
    q3 = simulate_variant(data, adjusted=True, variable_price=False, load_point=load_point, pv_point=pv_point, load_residual=load_residual, pv_residual=pv_residual, issued_residual=issued_residual, q2_cache=q2_cache, q3_cache=q3_cache, alpha=args.alpha)
    q4_2 = simulate_variant(data, adjusted=False, variable_price=True, load_point=load_point, pv_point=pv_point, load_residual=load_residual, pv_residual=pv_residual, issued_residual=issued_residual, q2_cache=q2_cache, q3_cache=q3_cache, alpha=args.alpha)
    q4_3 = simulate_variant(data, adjusted=True, variable_price=True, load_point=load_point, pv_point=pv_point, load_residual=load_residual, pv_residual=pv_residual, issued_residual=issued_residual, q2_cache=q2_cache, q3_cache=q3_cache, alpha=args.alpha)
    attach_variant_metadata(q2, variable_price=False)
    attach_variant_metadata(q3, variable_price=False)
    attach_variant_metadata(q4_2, variable_price=True)
    attach_variant_metadata(q4_3, variable_price=True)

    templates = root / "data" / "附件" / "附件5"
    write_result1(data, q1, templates / "result1.xlsx", results_dir / "result1.xlsx")
    write_annual_workbook(data, q2, templates / "result2.xlsx", results_dir / "result2.xlsx", adjusted=False)
    write_annual_workbook(data, q3, templates / "result3.xlsx", results_dir / "result3.xlsx", adjusted=True)
    write_annual_workbook(data, q4_2, templates / "result4-2.xlsx", results_dir / "result4-2.xlsx", adjusted=False)
    write_annual_workbook(data, q4_3, templates / "result4-3.xlsx", results_dir / "result4-3.xlsx", adjusted=True)

    q1_schedule, q1_aggregate = q1_frames(data, q1)
    q1_schedule.to_csv(results_dir / "问题1_逐时结果.csv", index=False, encoding="utf-8-sig")
    q1_aggregate.to_csv(results_dir / "问题1_四小时汇总.csv", index=False, encoding="utf-8-sig")
    variants = {"q2": (q2, False), "q3": (q3, True), "q4_2": (q4_2, False), "q4_3": (q4_3, True)}
    daily_frames = {}
    for name, (variant_results, adjusted) in variants.items():
        daily = daily_summary_frame(variant_results, adjusted=adjusted)
        daily.to_csv(results_dir / f"{name}_每日汇总.csv", index=False, encoding="utf-8-sig")
        daily_frames[name] = daily
        if not args.skip_csv:
            template_path = templates / ({"q2": "result2.xlsx", "q3": "result3.xlsx", "q4_2": "result4-2.xlsx", "q4_3": "result4-3.xlsx"}[name])
            wb = load_workbook(template_path, read_only=True, data_only=True)
            labels = [wb["计划购电量"].cell(1, col).value for col in range(2, 2 + N_SLOTS)]
            for r in variant_results:
                r["template_labels"] = labels
            interval_frame(data, variant_results, adjusted=adjusted).to_csv(
                results_dir / f"{name}_逐时结果.csv", index=False, encoding="utf-8-sig"
            )

    fmetrics = forecast_diagnostics(data, load_point, pv_point, load_residual, pv_residual, issued_residual)
    fmetrics.to_csv(results_dir / "预测误差指标.csv", index=False, encoding="utf-8-sig")
    raw_data_profile(data).to_csv(results_dir / "原始数据剖析.csv", index=False, encoding="utf-8-sig")
    write_representative_tables(root, variants, REPRESENTATIVE_DATES)
    rep_rows = []
    for name, (variant_results, adjusted) in variants.items():
        by_date = {r["date"].normalize(): r for r in variant_results}
        for date in REPRESENTATIVE_DATES:
            rep_rows.append({"variant": name, "date": date.date().isoformat(), **summarize_day(by_date[date], adjusted=adjusted)})
    pd.DataFrame(rep_rows).to_csv(results_dir / "典型日汇总.csv", index=False, encoding="utf-8-sig")

    q1_metrics = {
        "purchase_kwh": float(q1["grid_kwh"].sum()),
        "cash_cost_yuan": float(q1["cash_contract_cost"]),
        "charge_kwh": float(q1["charge_kwh"].sum()),
        "discharge_kwh": float(q1["discharge_kwh"].sum()),
        "pv_curtailed_kwh": float((data.typical_pv_kw * DT_HOURS - q1["pv_used_kwh"]).sum()),
        "soc_min_kwh": float(q1["soc_kwh"].min()),
        "soc_max_kwh": float(q1["soc_kwh"].max()),
        "max_balance_residual_kwh": float(q1["max_eq_residual"]),
    }
    annual = {}
    for name, daily in daily_frames.items():
        annual[name] = {
            "days": int(len(daily)),
            "plan_kwh": float(daily["plan_kwh"].sum()),
            "final_contract_kwh": float(daily["final_contract_kwh"].sum()),
            "cash_cost_yuan": float(daily["cash_cost_yuan"].sum()),
            "alternative_cash_cost_yuan": float(daily["alternative_cash_cost_yuan"].sum()),
            "emergency_kwh": float(daily["emergency_kwh"].sum()),
            "unused_contract_kwh": float(daily["unused_contract_kwh"].sum()),
            "pv_curtailed_kwh": float(daily["pv_curtailed_kwh"].sum()),
            "soc_end_kwh": float(daily.iloc[-1]["soc_end_kwh"]),
            "max_balance_residual_kwh": float(daily["max_balance_residual_kwh"].max()),
        }
    metrics = {
        "seed": SEED,
        "alpha": args.alpha,
        "time_mapping": "input timestamp is natural 10-minute interval endpoint; official template header retained by position",
        "q1": q1_metrics,
        "annual_feb_dec": annual,
        "comparisons": {
            "fixed_price_adjustment_saving_yuan": annual["q2"]["cash_cost_yuan"] - annual["q3"]["cash_cost_yuan"],
            "variable_price_adjustment_saving_yuan": annual["q4_2"]["cash_cost_yuan"] - annual["q4_3"]["cash_cost_yuan"],
            "q3_primary_vs_alternative_settlement_yuan": annual["q3"]["alternative_cash_cost_yuan"] - annual["q3"]["cash_cost_yuan"],
            "q4_3_primary_vs_alternative_settlement_yuan": annual["q4_3"]["alternative_cash_cost_yuan"] - annual["q4_3"]["cash_cost_yuan"],
        },
        "runtime_seconds": time.time() - started,
    }
    (results_dir / "核心指标.json").write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")

    inputs = [root / "data" / "C题.pdf"] + [root / "data" / "附件" / f"附件{i}.xlsx" for i in range(1, 5)]
    manifest = {
        "command": "python run_all.py --project-root .",
        "python": sys.version,
        "platform": platform.platform(),
        "packages": {"numpy": np.__version__, "pandas": pd.__version__, "scipy": scipy.__version__, "sklearn": sklearn.__version__},
        "seed": SEED,
        "parameters": {"alpha": args.alpha, "eta_charge": 0.9, "eta_discharge": 0.9, "soc_min_kwh": 1200, "soc_max_kwh": 10800, "power_limit_kw": 5000},
        "input_files": [file_metadata(p) for p in inputs],
        "hash_checks": "omitted per explicit user request; no hashes calculated or compared",
        "outputs": [str(p.relative_to(root)) for p in sorted(results_dir.glob("*"))],
        "runtime_seconds": metrics["runtime_seconds"],
    }
    manifest["command"] = f"python run_all.py --project-root . --alpha {args.alpha}"
    manifest["source_files"] = [file_metadata(p) for p in (root/"microgrid_core.py", root/"run_all.py", root/"evaluate_results.py")]
    manifest["validation"] = "every simulated day checked for physical feasibility, LP residuals, simultaneous flow and SOC continuity"
    (results_dir / "主运行记录.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(metrics, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
