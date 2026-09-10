"""Numerical/structural audit of the five official result workbooks, no hashes."""
from __future__ import annotations
import argparse
import json
from pathlib import Path
import numpy as np
import pandas as pd
from numpy.testing import assert_allclose
from openpyxl import load_workbook


def audit(root):
    result=root/"results"
    templates=root/"data"/"附件"/"附件5"
    logs=[]
    for filename in ("result1.xlsx","result2.xlsx","result3.xlsx","result4-2.xlsx","result4-3.xlsx"):
        wb=load_workbook(result/filename,data_only=False,read_only=True)
        template=load_workbook(templates/filename,data_only=False,read_only=True)
        assert wb.sheetnames==template.sheetnames,filename
        cells=formulas=0
        for ws in wb:
            for row in ws:
                for cell in row:
                    if cell.value is not None: cells+=1
                    assert cell.data_type!="e",(filename,ws.title,cell.coordinate)
                    formulas+=int(cell.data_type=="f")
            assert list(next(ws.values))==list(next(template[ws.title].values)),(filename,ws.title,"header")
        # Current templates and outputs are numeric; future formulas require recalc.
        assert formulas==0,(filename,"formula recalculation required")
        if filename=="result1.xlsx":
            csv=pd.read_csv(result/"问题1_逐时结果.csv")
            values=list(wb["计划购电量"].values)
            assert_allclose([r[1] for r in values[1:145]],csv.plan_purchase_kwh,atol=1e-7)
            groups=list(wb["充放电量"].values)
            for g in range(6):
                cut=csv.iloc[g*24:(g+1)*24]
                assert_allclose(groups[g+1][1:3],[cut.charge_kwh.sum(),cut.discharge_kwh.sum()],atol=1e-7)
            assert_allclose([groups[1][4],groups[2][4]],[6000,6000],atol=1e-7)
        else:
            name={"result2.xlsx":"q2","result3.xlsx":"q3","result4-2.xlsx":"q4_2","result4-3.xlsx":"q4_3"}[filename]
            adjusted=name in ("q3","q4_3")
            csv=pd.read_csv(result/f"{name}_逐时结果.csv",parse_dates=["date"])
            daily=pd.read_csv(result/f"{name}_每日汇总.csv",parse_dates=["date"])
            dates=pd.date_range("2025-02-01","2025-12-31")
            assert len(csv)==334*144 and len(daily)==334
            assert_allclose(csv.soc_start_kwh.to_numpy()[1:],csv.soc_end_kwh.to_numpy()[:-1],atol=1e-7)
            assert np.isfinite(csv.select_dtypes("number").to_numpy()).all()
            assert_allclose(csv.cash_cost_yuan,csv.settlement_cost_yuan+csv.emergency_cost_yuan,atol=1e-7)
            p=csv.initial_plan_kwh.to_numpy(); q=csv.final_contract_kwh.to_numpy(); c=csv.price_yuan_per_kwh.to_numpy()
            settlement=c*np.minimum(p,q)+1.5*c*np.maximum(q-p,0)+.5*c*np.maximum(p-q,0) if adjusted else c*p
            assert_allclose(csv.settlement_cost_yuan,settlement,atol=1e-7)
            assert_allclose(csv.emergency_cost_yuan,5*c*csv.emergency_kwh,atol=1e-7)
            for sheet,field in (("计划购电量","initial_plan_kwh"),)+((("调整购电量","final_contract_kwh"),) if adjusted else ()):
                values=list(wb[sheet].values)
                nonempty=[r for r in values[1:] if r[0] is not None]
                assert len(nonempty)==334,(filename,sheet,len(nonempty))
                assert pd.DatetimeIndex([r[0] for r in nonempty]).equals(dates)
                matrix=np.asarray([r[1:145] for r in nonempty],dtype=float)
                assert_allclose(matrix,csv[field].to_numpy().reshape(334,144),atol=1e-7)
                assert_allclose([r[145] for r in nonempty],matrix.sum(axis=1),atol=1e-7)
                costs=(c*p).reshape(334,144).sum(axis=1) if adjusted and sheet=="计划购电量" else csv.cash_cost_yuan.to_numpy().reshape(334,144).sum(axis=1)
                assert_allclose([r[146] for r in nonempty],costs,atol=1e-6)
            groups=list(wb["充放电量"].values)[1:]
            groups=[r for r in groups if r[1] is not None]
            assert len(groups)==334*6
            for d,date in enumerate(dates):
                cut=csv.iloc[d*144:(d+1)*144]
                assert pd.Timestamp(groups[d*6][0])==date
                for g in range(6):
                    row=groups[d*6+g]
                    assert row[1]==f"{g*4}:00-{(g+1)*4}:00"
                    part=cut.iloc[g*24:(g+1)*24]
                    assert_allclose(row[2:4],[part.charge_kwh.sum(),part.discharge_kwh.sum()],atol=1e-7)
                assert_allclose([groups[d*6][5],groups[d*6+1][5]],[cut.iloc[0].soc_start_kwh,cut.iloc[-1].soc_end_kwh],atol=1e-7)
            emergency=[r for r in list(wb["紧急购电量"].values)[1:] if r[0] is not None]
            expected=csv[csv.emergency_kwh>1e-8]
            if len(expected):
                assert len(emergency)==len(expected)
                assert pd.DatetimeIndex([r[0] for r in emergency]).equals(pd.DatetimeIndex(expected.date))
                assert [r[1] for r in emergency]==expected.natural_interval.tolist()
                assert_allclose([r[2] for r in emergency],expected.emergency_kwh,atol=1e-7)
            else:
                assert len(emergency)==1 and emergency[0][1:3]==("无",0)
            assert_allclose(daily.cash_cost_yuan,csv.cash_cost_yuan.to_numpy().reshape(334,144).sum(axis=1),atol=1e-6)
        logs.append(dict(file=filename,filled_cells=cells,formula_count=formulas,status="PASS"))
        wb.close();template.close()
    report=dict(status="author_numeric_audit_passed_not_independent_P2",checks=logs,hash_checks="omitted per user request")
    (result/"工作簿审计.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=="__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("--project-root",default=".")
    audit(Path(parser.parse_args().project_root).resolve())
