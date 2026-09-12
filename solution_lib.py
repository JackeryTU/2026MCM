# -*- coding: utf-8 -*-
"""编程手求解层公共工具：预测器、安全轨迹、年度回放、结果写出。

与 ``microgrid_core.py``（模型/优化内核）分离：本模块只做“把内核串成年度闭环”
与“把结果写成 5 份模板副本 + CSV + 图件”的工作，便于问题 2/3/4 复用。
"""
from __future__ import annotations

import datetime as _dt
import time

import numpy as np
import openpyxl
import pandas as pd

import microgrid_core as mc

ROOT = mc.ROOT
TPL_DIR = ROOT / "data" / "附件" / "附件5"
RES = ROOT / "results"
CSV = RES / "csv"
FIG = ROOT / "figures"

REP_DATES = [_dt.date(2025, 3, 20), _dt.date(2025, 6, 21),
             _dt.date(2025, 9, 23), _dt.date(2025, 12, 21)]
REP_SLOTS = [60, 72, 84, 96, 108, 120]        # 10:00,12:00,...,20:00 起（0 基自然索引）
DAY_INDEX = {d: (d - mc.DAY0).days for d in REP_DATES}
BLOCK_LABELS = ["0:00-4:00", "4:00-8:00", "8:00-12:00",
                "12:00-16:00", "16:00-20:00", "20:00-24:00"]


# ---------------------------------------------------------------------------
# 预测器
# ---------------------------------------------------------------------------
def predict_center(kind: str, d: int, data: dict):
    """中心预测（kW）。kind ∈ {F0, F1, F2}。"""
    load_act, pv_act = data["load_act"], data["pv_act"]
    if kind == "F1":
        return mc.forecast_f1(d, load_act, pv_act, data["load1"], data["pv1"])
    if kind == "F2":
        cand = [load_act[d - k] for k in (7, 14, 21) if d - k >= 0]
        if not cand:
            lhat = load_act[:d].mean(axis=0) if d >= 1 else data["load1"]
        elif len(cand) == 1:
            lhat = cand[0].copy()
        else:
            lhat = np.median(np.vstack(cand), axis=0)
        if d >= 1:
            vhat = pv_act[max(0, d - 7):d].mean(axis=0)
        else:
            vhat = data["pv1"].copy()
        return lhat, vhat
    if kind == "F0":
        if d == 0:
            return data["load1"].copy(), data["pv1"].copy()
        j = np.arange(max(0, d - 90), d)
        w = np.exp(-(d - j) / 60.0)
        dow = np.array([(mc.DAY0 + _dt.timedelta(days=int(x))).weekday() for x in j])
        wl = w * np.where(dow == (mc.DAY0 + _dt.timedelta(days=d)).weekday(), 2.0, 1.0)
        lhat = (load_act[j] * wl[:, None]).sum(axis=0) / wl.sum()
        vhat = (pv_act[j] * w[:, None]).sum(axis=0) / w.sum()
        rho = min(1.0, len(j) / 21.0)
        lhat = rho * lhat + (1 - rho) * data["load1"]
        vhat = rho * vhat + (1 - rho) * data["pv1"]
        return lhat, vhat
    raise ValueError(kind)


def net_residual_matrix(kind: str, data: dict) -> np.ndarray:
    """逐日净负荷残差 e^N = (L-V) - (Lhat-Vhat)，kW。全 365 天可算。"""
    load_act, pv_act = data["load_act"], data["pv_act"]
    out = np.zeros((365, mc.T))
    for d in range(365):
        lhat, vhat = predict_center(kind, d, data)
        out[d] = (load_act[d] - pv_act[d]) - (lhat - vhat)
    return out


def safe_trajectory(d: int, lhat_kw, vhat_kw, resid: np.ndarray, alpha: float, W: int):
    """日前安全轨迹（kW）：净负荷分位 + 保留弃光的非负重构。"""
    lo = max(1, d - W)
    if d <= 0 or lo >= d:
        q = np.zeros(mc.T)
    else:
        srt = np.sort(resid[lo:d], axis=0)
        k = int(np.ceil(alpha * srt.shape[0]))
        k = min(max(k, 1), srt.shape[0])
        q = srt[k - 1]
    ntilde = lhat_kw - vhat_kw + q
    vtilde = np.maximum(np.maximum(vhat_kw, -ntilde), 0.0)
    ltilde = vtilde + ntilde
    return ltilde, vtilde, ntilde


def point_error_metrics(kind: str, data: dict) -> dict:
    """评价期（2025-02-01—12-31）点预测误差，kW。"""
    eL, eV, eN = [], [], []
    for d in range(31, 365):
        lhat, vhat = predict_center(kind, d, data)
        eL.append(data["load_act"][d] - lhat)
        eV.append(data["pv_act"][d] - vhat)
        eN.append((data["load_act"][d] - data["pv_act"][d]) - (lhat - vhat))
    eL, eV, eN = np.asarray(eL), np.asarray(eV), np.asarray(eN)
    return {
        "load_MAE": float(np.abs(eL).mean()),
        "load_Bias": float(eL.mean()),
        "pv_MAE": float(np.abs(eV).mean()),
        "pv_Bias": float(eV.mean()),
        "net_MAE": float(np.abs(eN).mean()),
        "net_RMSE": float(np.sqrt((eN ** 2).mean())),
        "net_Bias": float(eN.mean()),
    }


# ---------------------------------------------------------------------------
# 年度回放
# ---------------------------------------------------------------------------
def run_year_q2(data: dict, kind: str = "F1", alpha: float = 0.9, W: int = 30,
                price_mode: str = "fixed", resid: np.ndarray | None = None,
                no_battery: bool = False, horizon: int = 36,
                n_days: int = 365, executor: str = "mpc",
                terminal_value_factor: float = 1.0) -> dict:
    """问题二闭环：安全轨迹 → 合同 LP → 滚动 MPC 执行 → 现金结算。

    MPC 与问题三、四共用 ``run_year_q34`` 的执行口径。首槽动作仅使用
    当时可用的预测；真实负荷和光伏只进入执行后的可行性裁剪，不反馈到
    首槽优化输入。从 2025-01-01 0:00 连续传递 SOC，1 月为预运行。
    """
    if not no_battery:
        if price_mode == "fixed":
            mpc_price_mode, settle_price = "fixed", "fixed"
        elif price_mode == "actual":
            mpc_price_mode, settle_price = "oracle", "actual"
        else:
            raise ValueError(f"未知价格模式：{price_mode}")
        run = run_year_q34(
            data, prof=None, kind=kind, alpha=alpha, W=W, mode="uniform",
            H=horizon, allowed=(0,), updates=(), pv_source="center",
            price_mode=mpc_price_mode, settle_price=settle_price,
            n_days=n_days, resid_override=resid, executor=executor,
            terminal_value_factor=terminal_value_factor,
        )
        return run

    load_act, pv_act = data["load_act"], data["pv_act"]
    s = float(mc.S0)
    n = n_days
    P = np.zeros((n, mc.T))        # 合同购电量 p
    X = np.zeros((n, mc.T))
    Y = np.zeros((n, mc.T))
    U = np.zeros((n, mc.T))
    Wd = np.zeros((n, mc.T))
    Vv = np.zeros((n, mc.T))
    S = np.zeros((n, mc.T))
    cash = np.zeros(n)
    pv_used_ratio = np.zeros(n)
    load_kwh = load_act * mc.DT
    pv_kwh = pv_act * mc.DT
    for d in range(n):
        c = data["price1"] if price_mode == "fixed" else data["price_act"][d]
        L = load_kwh[d]
        V = pv_kwh[d]
        p = np.maximum(L - V, 0.0)
        S[d] = s
        Vv[d] = np.minimum(V, L)
        U[d] = np.maximum(L - V - p, 0.0)
        Wd[d] = np.maximum(p - (L - Vv[d]), 0.0)
        cash[d] = float(np.sum(c * p) + mc.EMERGENCY_FACTOR * np.sum(c * U[d]))
        P[d] = p
    return dict(P=P, X=X, Y=Y, U=U, W=Wd, V=Vv, S=S, cash=cash,
                kind=kind, alpha=alpha, resid_W=W, price_mode=price_mode,
                S_init=mc.S0, S_end=float(s), pv_curtail_ratio=pv_used_ratio,
                resid=resid, load_kwh=load_kwh, pv_kwh=pv_kwh,
                load_kw=load_act, pv_kw=pv_act)


def perfect_foresight_q2(data: dict, horizon: int = 36) -> dict:
    """理想信息参照：真实轨迹制定合同，再用同一滚动 MPC 执行。"""
    load_act, pv_act = data["load_act"], data["pv_act"]
    s = float(mc.S0)
    n = 365
    out = {k: np.zeros((n, mc.T)) for k in ("P", "X", "Y", "U", "W", "V", "S")}
    cash = np.zeros(n)
    for d in range(n):
        c = data["price1"]
        L = load_act[d] * mc.DT
        V = pv_act[d] * mc.DT
        lp = mc.solve_day_lp(c, L, V, s, lam=mc.ETA_D * float(np.median(c)))
        p = np.maximum(lp["g"], 0.0)
        pv_versions = np.tile(V, (N_ISSUE, 1))
        price_seg_fn = lambda t, K, c=c: c[t:K + 1]
        ex = rolling_execute_q34(
            p, L, pv_versions, L, V, price_seg_fn, s,
            horizon=horizon, allowed=(0,), time_limit=5.0,
        )
        out["P"][d] = p
        out["X"][d], out["Y"][d] = ex["x"], ex["y"]
        out["U"][d], out["W"][d], out["V"][d], out["S"][d] = ex["u"], ex["w"], ex["v"], ex["S"]
        cash[d] = mc.settle_q12(out["P"][d], out["U"][d], c)
        s = float(min(max(ex["S"][-1], mc.S_MIN), mc.S_MAX))
    out["cash"] = cash
    out["S_end"] = s
    return out


# ---------------------------------------------------------------------------
# 指标与表格
# ---------------------------------------------------------------------------
def official_slice(arr, offset: int = 31):
    return arr[offset:365]


def year_metrics(run: dict, offset: int = 31) -> dict:
    cash = official_slice(run["cash"], offset)
    U = official_slice(run["U"], offset)
    P = official_slice(run["P"], offset)
    Y = official_slice(run["Y"], offset)
    X = official_slice(run["X"], offset)
    Wd = official_slice(run["W"], offset)
    S = official_slice(run["S"], offset)
    price = None
    return {
        "daily_cash": cash,
        "cash_total": float(cash.sum()),
        "contract_kwh": float(P.sum()),
        "charge_kwh": float(X.sum()),
        "discharge_kwh": float(Y.sum()),
        "unused_contract_kwh": float(Wd.sum()),
        "emergency_kwh": float(U.sum()),
        "emergency_slots": int((U > 1e-8).sum()),
        "S_min": float(S.min()), "S_max": float(S.max()),
        "S_end": float(run["S_end"]),
    }


def energy_block_table(run: dict, day_indices, offset: int = 31) -> pd.DataFrame:
    """表 2：重点日 6 个 4 h 段的实际充放电量与日初/日末 SOC。"""
    rows = []
    for d in day_indices:
        ch = mc.four_hour_blocks(run["X"][d])
        dis = mc.four_hour_blocks(run["Y"][d])
        s0 = run["S"][d - 1][-1] if d >= 1 else mc.S0
        for k in range(6):
            rows.append({"日期": str(mc.date_of(d)), "时间段": BLOCK_LABELS[k],
                         "充电量_kWh": ch[k], "放电量_kWh": dis[k],
                         "日初储电量_kWh": s0, "日末储电量_kWh": run["S"][d][-1]})
    return pd.DataFrame(rows)


def slot_table(run: dict, price_day, day_indices, slots=None) -> pd.DataFrame:
    """表 1：重点日指定槽的合同购电、实际购电、应急与充放电。"""
    slots = REP_SLOTS if slots is None else slots
    rows = []
    for d in day_indices:
        c = price_day if np.ndim(price_day) == 1 else price_day[d]
        for t in slots:
            rows.append({
                "日期": str(mc.date_of(d)),
                "时段": f"{mc.slot_label_start(t)}-{mc.slot_label_end(t)}",
                "电价_元每kWh": c[t],
                "负荷_kW": float(run["load_kw"][d][t]),
                "光伏实际_kW": float(run["pv_kw"][d][t]),
                "合同购电_p_kWh": run["P"][d][t],
                "实际购电_kWh": run["P"][d][t] - run["W"][d][t] + run["U"][d][t],
                "应急购电_u_kWh": run["U"][d][t],
                "充电量_kWh": run["X"][d][t],
                "放电量_kWh": run["Y"][d][t],
                "槽末储电量_kWh": run["S"][d][t],
            })
    return pd.DataFrame(rows)


def emergency_intervals(U_day: np.ndarray, date, threshold: float = 1e-8,
                        max_intervals: int = 3) -> list:
    """同一日内连续 u>threshold 的最大区间（不跨午夜）。"""
    spans = mc.merge_intervals(np.asarray(U_day) > threshold)
    out = []
    for i0, i1 in spans[:max_intervals]:
        out.append({"日期": str(date), "购电时间段": mc.interval_label(i0, i1),
                    "购电量_kWh": float(U_day[i0:i1 + 1].sum())})
    return out


def emergency_table(run: dict, day_indices, threshold: float = 1e-8) -> pd.DataFrame:
    rows = []
    for d in day_indices:
        ivs = emergency_intervals(run["U"][d], mc.date_of(d), threshold)
        if not ivs:
            rows.append({"日期": str(mc.date_of(d)), "购电时间段": "无", "购电量_kWh": 0.0})
        rows.extend(ivs)
    return pd.DataFrame(rows)


def check_run(run: dict, offset: int = 31) -> dict:
    """问题二数值门禁；保留旧字段并补齐边界、SOC 链和结算检查。"""
    P, X, Y, U, Wd, V, S = (np.asarray(run[k])
                              for k in ("P", "X", "Y", "U", "W", "V", "S"))
    n = P.shape[0]
    L = np.asarray(run["load_kwh"])[:n]
    pv = np.asarray(run["pv_kwh"])[:n]
    bal = np.abs((P - Wd) + V + Y + U - X - L)
    soc_chain = 0.0
    for d in range(1, n):
        expected = S[d - 1, -1] + mc.ETA_C * X[d, 0] - Y[d, 0] / mc.ETA_D
        soc_chain = max(soc_chain, abs(S[d, 0] - expected))
    price = run.get("price_settle")
    settle = float("nan") if price is None else float(np.abs(
        np.asarray(run["cash"])[:n]
        - (np.asarray(run["cash_contract"])[:n]
           + mc.EMERGENCY_FACTOR * (np.asarray(price)[:n] * U).sum(axis=1))
    ).max())
    bounds = {
        "bound_p": float(-min(P.min(), 0.0)),
        "bound_x": float(max(-min(X.min(), 0.0), (X - mc.B_MAX).max())),
        "bound_y": float(max(-min(Y.min(), 0.0), (Y - mc.B_MAX).max())),
        "bound_v": float(max(-min(V.min(), 0.0), (V - pv).max())),
        "bound_u": float(-min(U.min(), 0.0)),
        "bound_w": float(max(-min(Wd.min(), 0.0), (Wd - P).max())),
        "bound_S": float(max(mc.S_MIN - S.min(), S.max() - mc.S_MAX)),
    }
    return {
        "balance_max": float(bal.max()),
        "overlap_max": float(np.minimum(X, Y).max()),
        "charge_emergency_max": float(np.minimum(X, U).max()),
        "S_min": float(S.min()), "S_max": float(S.max()),
        "soc_chain_max_kWh": float(soc_chain),
        "soc_end_mismatch_kWh": float(abs(S[-1, -1] - run["S_end"])),
        "settle_flow_max_yuan": settle,
        "official_days": int(max(0, n - offset)),
        **bounds,
    }


# ---------------------------------------------------------------------------
# 模板写出
# ---------------------------------------------------------------------------
def _date_blocks(ws, n_rows: int) -> list:
    """按模板第 1 列出现的日期切分数据块，返回 [(首行, 日期字符串), ...]。

    这样写回时不依赖固定行距，模板里的 ``⁝`` 占位行会被自然跳过。
    """
    blocks = []
    for row in range(2, n_rows + 1):
        value = ws.cell(row=row, column=1).value
        if value is None:
            continue
        text = str(value).strip()
        if text[:4].isdigit() and "-" in text:
            blocks.append((row, text[:10]))
    return blocks


def _day_index_of(date_text: str, offset: int = 31):
    try:
        d = _dt.date.fromisoformat(date_text[:10])
    except ValueError:
        return None
    idx = (d - mc.DAY0).days
    return idx if offset <= idx < 365 else None


def _fill_plan_sheet(ws, run: dict, offset: int = 31, key: str = "P",
                     fee_key: str = "cash") -> None:
    """计划/调整购电量 335×147：行序填 144 列 + 全天购电量 + 全天购电费。"""
    arr = run[key]
    fee = run.get(fee_key, run["cash"])
    for r, d in enumerate(range(offset, arr.shape[0])):
        row = 2 + r
        a = arr[d]
        for t in range(mc.T):
            ws.cell(row=row, column=2 + t).value = float(a[t])
        ws.cell(row=row, column=146).value = float(a.sum())
        ws.cell(row=row, column=147).value = float(fee[d])


def _fill_cycle_sheet(ws, run: dict, offset: int = 31) -> None:
    """充放电量：按模板自带的日期块填 (6 段充放电, 日初/日末储电量)。"""
    for base, date_text in _date_blocks(ws, ws.max_row):
        d = _day_index_of(date_text, offset)
        if d is None or d >= run["X"].shape[0]:
            continue
        ch = mc.four_hour_blocks(run["X"][d])
        dis = mc.four_hour_blocks(run["Y"][d])
        for k in range(6):
            ws.cell(row=base + k, column=3).value = float(ch[k])
            ws.cell(row=base + k, column=4).value = float(dis[k])
        ws.cell(row=base, column=6).value = float(
            run["S"][d - 1][-1] if d >= 1 else mc.S0)
        ws.cell(row=base + 1, column=6).value = float(run["S"][d][-1])


def _fill_emergency_sheet(ws, run: dict, offset: int = 31) -> None:
    """紧急购电量：按模板自带的日期块填 (购电时间段, 购电量)。"""
    for base, date_text in _date_blocks(ws, ws.max_row):
        d = _day_index_of(date_text, offset)
        if d is None or d >= run["U"].shape[0]:
            continue
        ivs = emergency_intervals(run["U"][d], mc.date_of(d))
        if not ivs:
            ws.cell(row=base, column=2).value = "无"
            ws.cell(row=base, column=3).value = 0.0
            continue
        for k, iv in enumerate(ivs):
            ws.cell(row=base + k, column=2).value = iv["购电时间段"]
            ws.cell(row=base + k, column=3).value = iv["购电量_kWh"]


def write_workbook(template: str, out_name: str, sheets: dict) -> None:
    """按模板复制后逐表填入，保留表名与列序。sheets: {sheet_name: callable(ws)}。"""
    RES.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(TPL_DIR / template)
    for name, filler in sheets.items():
        filler(wb[name])
    wb.save(RES / out_name)


def monthly_summary(run: dict, offset: int = 31) -> pd.DataFrame:
    rows = []
    for d in range(offset, run["P"].shape[0]):
        dt = mc.date_of(d)
        rows.append({"日期": dt, "月份": dt.strftime("%Y-%m"),
                     "现金费_元": run["cash"][d],
                     "应急购电量_kWh": run["U"][d].sum(),
                     "合同购电量_kWh": run["P"][d].sum(),
                     "充电量_kWh": run["X"][d].sum(),
                     "放电量_kWh": run["Y"][d].sum(),
                     "弃购量_kWh": run["W"][d].sum(),
                     "弃光量_kWh": float((run["pv_kwh"][d] - run["V"][d]).sum())
                     if "pv_kwh" in run else np.nan,
                     "日末储电量_kWh": run["S"][d][-1]})
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# 问题三/四：官方光伏预报档案、分层风险、滚动合同—执行闭环
# ---------------------------------------------------------------------------
ISSUE_SLOTS = (0, 36, 72, 108)          # 0:00 / 6:00 / 12:00 / 18:00 发布时刻
ISSUE_HOURS = (0, 6, 12, 18)
N_ISSUE = 4
N_STRATA = 6                             # 预见期分层：0–1/1–2/…/5–6 h


def build_pv_archive(data: dict, W: int = 30) -> np.ndarray:
    """逐日归档“当时实际发出”的附件3 校正+插值 10 min 预报，形状 (365, 4, 145)。

    偏差校正只用 j<d 最近 W 个已完成日的同 (发布 r, 提前期 h) 原始误差均值（因果）；
    零提前期锚点取发布时刻已完成的实际光伏端点；未获准的提前期不伪造。
    """
    pv_act = data["pv_act"]
    fc = data["fc_raw"]
    prof = np.zeros((365, N_ISSUE, mc.T + 1))
    for d in range(365):
        lo = max(0, d - W)
        for r in range(N_ISSUE):
            corr = np.zeros(24)
            if d > lo:
                for h in range(1, 25):
                    pos = ISSUE_SLOTS[r] + 6 * h
                    if pos > mc.T:
                        break
                    corr[h - 1] = float(np.mean(
                        [pv_act[j, pos - 1] - fc[j, r, h - 1] for j in range(lo, d)]))
            corrected = np.maximum(fc[d, r, :] + corr, 0.0)
            if r == 0:
                anchor = float(pv_act[d - 1, -1]) if d >= 1 else 0.0
            else:
                anchor = float(pv_act[d, ISSUE_SLOTS[r] - 1])
            prof[d, r] = mc.interp_official_profile(corrected, ISSUE_HOURS[r], anchor)
    return prof


def issue_index_of_slot(t: int, allowed=(0, 1, 2, 3)) -> int:
    """槽 t 生效的发布版本：允许集合中最后一个发布时刻 <= t 的版本。"""
    best = 0
    for r in allowed:
        if ISSUE_SLOTS[r] <= t:
            best = r
    return best


def version_of_slot_array(allowed=(0, 1, 2, 3)) -> np.ndarray:
    return np.array([issue_index_of_slot(t, allowed) for t in range(mc.T)], dtype=int)


def pv_version_curve(prof_day: np.ndarray, allowed=(0, 1, 2, 3)) -> np.ndarray:
    """按“未获准新预报就继续用最近获准版本”展开成 (144,) 逐槽预报 kW。

    对齐口径：附件数据列 = 该 10 min 区间的结束时刻，故区间 t 取插值曲线上
    的时刻点 t+1（= (t+1)*10 min），与附件 2 逐槽实际值的区间口径一致。
    """
    return np.array([prof_day[issue_index_of_slot(t, allowed), t + 1]
                     for t in range(mc.T)])


def horizon_stratum(allowed=(0, 1, 2, 3)) -> np.ndarray:
    """按实际允许的最近预报版本计算逐槽预见期分层 0..5。"""
    out = np.empty(mc.T, dtype=int)
    for t in range(mc.T):
        r = issue_index_of_slot(t, allowed)
        out[t] = min((t - ISSUE_SLOTS[r]) // 6, N_STRATA - 1)
    return out


STRATUM = horizon_stratum()


def q3_net_residual(data: dict, prof: np.ndarray, allowed=(0, 1, 2, 3),
                    kind: str = "F1", pv_source: str = "archive"):
    """逐日净负荷残差（kW）与生效的光伏预报曲线 (365,144)。

    pv_source='archive' 用遮罩后版本的官方光伏预报；'center' 用 F1/F2 中心光伏
    预测（对应 Q4-2 沿用问题二风险体系的写法）。
    """
    pv_ver = np.zeros((365, mc.T))
    resid = np.zeros((365, mc.T))
    for d in range(365):
        lhat, vhat = predict_center(kind, d, data)
        if pv_source == "archive":
            pv_ver[d] = pv_version_curve(prof[d], allowed)
        else:
            pv_ver[d] = vhat
        resid[d] = (data["load_act"][d] - data["pv_act"][d]) - (lhat - pv_ver[d])
    return resid, pv_ver


def q3_reserve(d: int, resid: np.ndarray, alpha: float, W: int,
               mode: str = "stratified", allowed=(0, 1, 2, 3)):
    """逐槽风险储备（kW）与分层尺度 sigma(1..6)。

    mode='stratified'：按预见期分层标准化后合并分位，再乘回单调尺度；
    mode='uniform'：与问题二一致的逐槽净负荷残差分位。
    """
    lo = max(1, d - W)
    if d <= 0 or lo >= d:
        return np.zeros(mc.T), np.ones(N_STRATA)
    hist = resid[lo:d]
    if mode == "uniform":
        return (np.array([mc.empirical_quantile(hist[:, t], alpha)
                          for t in range(mc.T)]), np.ones(N_STRATA))
    stratum = horizon_stratum(allowed)
    sig = np.ones(N_STRATA)
    for h in range(N_STRATA):
        cols = np.where(stratum == h)[0]
        if cols.size:
            sig[h] = float(np.sqrt((hist[:, cols] ** 2).mean()))
    sig = np.maximum.accumulate(np.maximum(sig, 1e-9))     # 单调非减的经验尺度
    z = hist / sig[stratum][None, :]
    qs = mc.empirical_quantile(z.ravel(), alpha)
    return sig[stratum] * qs, sig


def safe_from_reserve(lhat_kwh, vhat_kwh, reserve_kwh):
    """把中心预测与风险储备合成为非负可行的安全负荷/光伏轨迹（kWh）。"""
    ntilde = lhat_kwh - vhat_kwh + reserve_kwh
    vtilde = np.maximum(np.maximum(vhat_kwh, -ntilde), 0.0)
    return vtilde + ntilde, vtilde, ntilde


def rolling_execute_q34(g_contract, load_pred_kwh, pv_ver_kwh, load_kwh, pv_kwh,
                        price_seg_fn, s_init, horizon: int = 36,
                        allowed=(0, 1, 2, 3), time_limit: float = 5.0,
                        contract_update_fn=None, executor: str = "mpc",
                        terminal_value_factor: float = 1.0):
    """§7.2/§7.3 滚动 LP + 槽内裁剪执行。

    LP 只用当时获准的预测（load_pred 为日初负荷预测，pv_ver_kwh 为四个发布版本
    各自的 10 min 光伏预报）与当前合同；槽内用当槽真实值按 §7.3 裁剪。
    超时/非最优时回退到时序对齐的延迟贪心首槽动作。
    """
    n = mc.T
    contract = np.asarray(g_contract, dtype=float).copy()
    x = np.zeros(n); y = np.zeros(n); u = np.zeros(n)
    w = np.zeros(n); v = np.zeros(n); S = np.zeros(n)
    plan_x = np.zeros(n); plan_y = np.zeros(n); plan_s_next = np.zeros(n)
    s = float(s_init)
    stat = {"clip_charge": 0, "clip_discharge": 0, "fallback": 0,
            "updates": 0, "lp_time": []}
    if executor not in ("mpc", "greedy"):
        raise ValueError(f"未知执行器：{executor}")
    prev_net = None
    for t in range(n):
        if contract_update_fn is not None:
            q_rem = contract_update_fn(t, s, contract)
            if q_rem is not None:
                q_rem = np.asarray(q_rem, dtype=float)
                if q_rem.shape != (n - t,):
                    raise ValueError(f"合同更新长度异常：t={t}, shape={q_rem.shape}")
                contract[t:] = np.maximum(q_rem, 0.0)
                stat["updates"] += 1
        r = issue_index_of_slot(t, allowed)
        load_curve = (load_pred_kwh[r] if np.ndim(load_pred_kwh) == 2
                      else load_pred_kwh)
        if executor == "mpc":
            K = min(t + horizon - 1, n - 1)
            seg = slice(t, K + 1)
            c_seg = price_seg_fn(t, K)
            lam = terminal_value_factor * mc.ETA_D * float(np.median(c_seg))
            tic = time.perf_counter()
            act = mc.rolling_lp_action(s, contract[seg], load_curve[seg],
                                       pv_ver_kwh[r][seg], c_seg, lam,
                                       time_limit=time_limit)
            stat["lp_time"].append(time.perf_counter() - tic)
            if act is None:
                net_est = float(prev_net) if prev_net is not None else float(
                    load_curve[t] - pv_ver_kwh[r][t])
                if net_est < 0.0:
                    xm = min(mc.B_MAX, -net_est,
                             max((mc.S_MAX - s) / mc.ETA_C, 0.0))
                    ym = 0.0
                else:
                    xm = 0.0
                    ym = min(mc.B_MAX, net_est,
                             max((s - mc.S_MIN) * mc.ETA_D, 0.0))
                stat["fallback"] += 1
            else:
                xm, ym = act
        else:
            surplus = max(contract[t] + pv_ver_kwh[r][t] - load_curve[t], 0.0)
            deficit = max(load_curve[t] - contract[t] - pv_ver_kwh[r][t], 0.0)
            xm = min(mc.B_MAX, surplus, max((mc.S_MAX - s) / mc.ETA_C, 0.0))
            ym = min(mc.B_MAX, deficit, max((s - mc.S_MIN) * mc.ETA_D, 0.0))
        xe, ye, ue, we, ve, _ = mc.clip_slot(xm, ym, contract[t], load_kwh[t],
                                             pv_kwh[t])
        plan_x[t], plan_y[t] = xm, ym
        plan_s_next[t] = min(max(s + mc.ETA_C * xm - ym / mc.ETA_D,
                                 mc.S_MIN), mc.S_MAX)
        if xm - xe > 1e-9:
            stat["clip_charge"] += 1
        if ym - ye > 1e-9:
            stat["clip_discharge"] += 1
        x[t], y[t], u[t], w[t], v[t] = xe, ye, ue, we, ve
        s = s + mc.ETA_C * xe - ye / mc.ETA_D
        s = min(max(s, mc.S_MIN), mc.S_MAX)
        S[t] = s
        prev_net = float(load_kwh[t] - pv_kwh[t])
    return {"x": x, "y": y, "u": u, "w": w, "v": v, "S": S,
            "plan_x": plan_x, "plan_y": plan_y,
            "plan_s_next": plan_s_next,
            "contract": contract, "stats": stat}


def run_year_q34(data: dict, prof: np.ndarray | None = None, kind: str = "F1",
                 alpha: float = 0.9, W: int = 30, mode: str = "stratified",
                 H: int = 36, allowed=(0, 1, 2, 3), updates=(1, 2, 3),
                 pv_source: str = "archive", price_mode: str = "fixed",
                 price_base: str = "median30", settle_price: str = "fixed",
                 time_limit: float = 5.0, n_days: int = 365,
                 resid_override: np.ndarray | None = None,
                 executor: str = "mpc", terminal_value_factor: float = 1.0) -> dict:
    """问题三/四闭环：0 时合同 LP → 6/12/18 合同调整 → 滚动 LP 执行 → 结算。

    price_mode='fixed' 用附件1 典型日价；'forecast' 用附件4 的预测价（含日内修正）。
    settle_price='fixed' 按典型日价结算；'actual' 按附件4 真实价结算。
    """
    if prof is None and pv_source == "archive":
        prof = build_pv_archive(data, W)
    resid, pv_ver = q3_net_residual(data, prof, allowed, kind, pv_source)
    if resid_override is not None:
        resid = np.asarray(resid_override, dtype=float)
    load_kw = data["load_act"]
    pv_kw = data["pv_act"]
    load_kwh = load_kw * mc.DT
    pv_kwh = pv_kw * mc.DT
    keys = ("P", "Q", "X", "Y", "U", "W", "V", "S",
            "plan_x", "plan_y", "plan_s_next")
    arr = {k: np.zeros((n_days, mc.T)) for k in keys}
    cash = np.zeros(n_days)
    cash_plan = np.zeros(n_days)
    cash_contract = np.zeros(n_days)
    price_settle = np.zeros((n_days, mc.T))
    stats = {"clip_charge": 0, "clip_discharge": 0, "fallback": 0,
             "updates": 0, "lp_time": []}
    # 分日计数：年度指标只统计正式窗口（offset 之后），1 月预运行不计入
    daily_cnt = {k: np.zeros(n_days, dtype=int)
                 for k in ("clip_charge", "clip_discharge", "fallback",
                           "updates")}
    lp_by_day = []
    s = float(mc.S0)
    for d in range(n_days):
        lhat_kw, vhat_kw = predict_center(kind, d, data)
        load_pred = lhat_kw * mc.DT
        if pv_source == "archive":
            # 区间 t 的预报取插值曲线时刻点 t+1（区间结束时刻口径）
            pv_curves = np.vstack([prof[d, r, 1:mc.T + 1] for r in range(N_ISSUE)]) * mc.DT
        else:
            pv_curves = np.tile(vhat_kw * mc.DT, (N_ISSUE, 1))
        if price_mode == "fixed":
            forecast_price = data["price1"]
            price_seg_fn = (lambda t, K, bp=forecast_price: bp[t:K + 1])
        elif price_mode == "oracle":
            # 附加理想信息基准：0:00 即知全天真实价格，其余不变（Q4 专用）
            forecast_price = data["price_act"][d]
            price_seg_fn = (lambda t, K, bp=forecast_price: bp[t:K + 1])
        else:
            forecast_price = mc.price_forecast_0h(d, data["price_act"], data["price1"],
                                                   price_base)

            def price_seg_fn(t, K, bp=forecast_price, d=d):
                return mc.price_intraday_correct(bp, data["price_act"][d], t)[t:K + 1]

        settle_day_price = data["price1"] if settle_price == "fixed" else data["price_act"][d]
        price_settle[d] = settle_day_price
        # 风险储备在 kW 空间估计，转成与安全轨迹同量纲的 kWh 再合成
        reserve_kw, _ = q3_reserve(d, resid, alpha, W, mode, allowed)
        reserve = reserve_kw * mc.DT
        # 合同层与执行层共用同一风险修正预测，避免 MPC 把安全余量丢掉。
        safe_load_curves = np.zeros_like(pv_curves)
        safe_pv_curves = np.zeros_like(pv_curves)
        for r in range(N_ISSUE):
            safe_load_curves[r], safe_pv_curves[r], _ = safe_from_reserve(
                load_pred, pv_curves[r], reserve
            )
        # 0:00 合同 p = q(0)：只有 0 时发布的版本可用
        ltil, vtil = safe_load_curves[0], safe_pv_curves[0]
        lp = mc.solve_day_lp(forecast_price, ltil, vtil, s,
                             lam=terminal_value_factor * mc.ETA_D * float(np.median(forecast_price)))
        p = np.maximum(lp["g"], 0.0)
        q = p.copy()
        cash_plan[d] = float(np.sum(forecast_price * p))
        update_at = {ISSUE_SLOTS[r]: r for r in updates}

        def update_contract(t, current_s, current_q):
            r_up = update_at.get(t)
            if r_up is None:
                return None
            c_up = price_seg_fn(t, mc.T - 1)
            ltil_up = safe_load_curves[r_up][t:]
            vtil_up = safe_pv_curves[r_up][t:]
            # 最终结算始终相对 0:00 原计划 p，因此每次重算也必须以 p 为基准。
            up = mc.contract_update_lp(
                p[t:], c_up, ltil_up, vtil_up, current_s,
                lam=terminal_value_factor * mc.ETA_D * float(np.median(c_up)),
            )
            return None if up is None else up["q"]

        ex = rolling_execute_q34(q, safe_load_curves, safe_pv_curves,
                                 load_kwh[d], pv_kwh[d],
                                 price_seg_fn, s, horizon=H, allowed=allowed,
                                 time_limit=time_limit,
                                 contract_update_fn=update_contract if updates else None,
                                 executor=executor,
                                 terminal_value_factor=terminal_value_factor)
        q = ex["contract"]
        cash_contract[d] = mc.settle_q3(p, q, np.zeros(mc.T), settle_day_price)
        for k, key in (("x", "X"), ("y", "Y"), ("u", "U"), ("w", "W"),
                       ("v", "V"), ("S", "S"), ("plan_x", "plan_x"),
                       ("plan_y", "plan_y"), ("plan_s_next", "plan_s_next")):
            arr[key][d] = ex[k]
        arr["P"][d] = p
        arr["Q"][d] = q
        cash[d] = mc.settle_q3(p, q, ex["u"], settle_day_price)
        for k in ("clip_charge", "clip_discharge", "fallback", "updates"):
            stats[k] += ex["stats"][k]
        stats["lp_time"].extend(ex["stats"]["lp_time"])
        for k in ("clip_charge", "clip_discharge", "fallback", "updates"):
            daily_cnt[k][d] += ex["stats"][k]
        lp_by_day.append(list(ex["stats"]["lp_time"]))
        s = float(min(max(ex["S"][-1], mc.S_MIN), mc.S_MAX))
    stats["daily_cnt"] = daily_cnt
    stats["lp_by_day"] = lp_by_day
    out = dict(arr)
    out.update(cash=cash, cash_plan=cash_plan, cash_contract=cash_contract,
               price_settle=price_settle,
               stats=stats, kind=kind, alpha=alpha, resid_W=W, mode=mode, H=H,
               allowed=tuple(allowed), updates=tuple(updates), pv_source=pv_source,
               price_mode=price_mode, price_base=price_base,
               settle_price=settle_price, executor=executor,
               terminal_value_factor=terminal_value_factor,
               S_init=mc.S0, S_end=float(s),
               resid=resid, pv_ver=pv_ver, load_kwh=load_kwh, pv_kwh=pv_kwh,
               load_kw=load_kw, pv_kw=pv_kw)
    return out


def _sliced_count(stats: dict, key: str, offset: int) -> int:
    """统计正式窗口内的发生次数；缺分日计数时退回全窗口累计。"""
    dc = stats.get("daily_cnt")
    if dc is None:
        return int(stats[key])
    return int(np.asarray(dc[key])[offset:].sum())


def year_metrics_q34(run: dict, offset: int = 31) -> dict:
    """问题三/四的年度指标（含合同调整额、裁剪与耗时统计）。"""
    P = official_slice(run["P"], offset)
    Q = official_slice(run["Q"], offset)
    U = official_slice(run["U"], offset)
    X = official_slice(run["X"], offset)
    Y = official_slice(run["Y"], offset)
    Wd = official_slice(run["W"], offset)
    S = official_slice(run["S"], offset)
    lp = np.asarray(run["stats"]["lp_time"], dtype=float)
    by_day = run["stats"].get("lp_by_day")
    if by_day is not None:
        # lp_time 逐次记录覆盖全年，正式窗口的耗时统计按日切片重建
        lp = np.asarray([v for day in by_day[offset:] for v in day], dtype=float)
    return {
        "daily_cash": official_slice(run["cash"], offset),
        "cash_total": float(official_slice(run["cash"], offset).sum()),
        "plan_kwh": float(P.sum()),
        "contract_kwh": float(Q.sum()),
        "adjust_up_kwh": float(np.maximum(Q - P, 0.0).sum()),
        "adjust_down_kwh": float(np.maximum(P - Q, 0.0).sum()),
        "charge_kwh": float(X.sum()),
        "discharge_kwh": float(Y.sum()),
        "unused_contract_kwh": float(Wd.sum()),
        "emergency_kwh": float(U.sum()),
        "emergency_slots": int((U > 1e-8).sum()),
        "S_min": float(S.min()), "S_max": float(S.max()), "S_end": float(run["S_end"]),
        "updates": _sliced_count(run["stats"], "updates", offset),
        "clip_charge": _sliced_count(run["stats"], "clip_charge", offset),
        "clip_discharge": _sliced_count(run["stats"], "clip_discharge", offset),
        "fallback": _sliced_count(run["stats"], "fallback", offset),
        "lp_ms_median": float(np.median(lp) * 1e3) if lp.size else float("nan"),
        "lp_ms_p95": float(np.percentile(lp, 95) * 1e3) if lp.size else float("nan"),
        "lp_ms_max": float(lp.max() * 1e3) if lp.size else float("nan"),
    }


def check_run_q34(run: dict, offset: int = 31) -> dict:
    """问题三/四数值与因果门禁（§11.4 与 §7.3）。

    覆盖：母线平衡、SOC/功率/非负边界、充放互斥、应急不主动充电、跨日 SOC
    连续、日末状态一致、结算流水与汇总一致。返回各项最大偏差。
    """
    P, Q, X, Y, U, Wd, V, S = (np.asarray(run[k])
                              for k in ("P", "Q", "X", "Y", "U", "W", "V", "S"))
    n = P.shape[0]
    # 短运行（冒烟）时全长实际序列需要截到与计划同长，避免广播错位
    L, pv = np.asarray(run["load_kwh"])[:n], np.asarray(run["pv_kwh"])[:n]
    bal = np.abs((Q - Wd) + V + Y + U - X - L)
    # 边界：报告最大越界量（负值截断为 0）
    b_p = float(-min(P.min(), 0.0))
    b_q = float(-min(Q.min(), 0.0))
    b_x = float(max(-min(X.min(), 0.0), (X - mc.B_MAX).max()))
    b_y = float(max(-min(Y.min(), 0.0), (Y - mc.B_MAX).max()))
    b_v = float(max(-min(V.min(), 0.0), (V - pv).max()))
    b_u = float(-min(U.min(), 0.0))
    b_w = float(max(-min(Wd.min(), 0.0), (Wd - Q).max()))
    b_s = float(max(mc.S_MIN - S.min(), S.max() - mc.S_MAX))
    chain = 0.0
    for d in range(1, n):
        chain = max(chain, abs((S[d][0] - S[d - 1][-1])
                              - (mc.ETA_C * X[d][0] - Y[d][0] / mc.ETA_D)))
    # 结算流水与汇总一致：cash = cash_contract + 5 * Σ c_s u
    price = run.get("price_settle")
    if price is None:
        flow = float("nan")
    else:
        flow = float(np.abs(run["cash"] - (run["cash_contract"]
                                          + mc.EMERGENCY_FACTOR * (price * U).sum(axis=1))).max())
    return {
        "balance_max_kWh": float(bal.max()),
        "overlap_max_kWh": float(np.minimum(X, Y).max()),
        "charge_emergency_max_kWh": float(np.minimum(X, U).max()),
        "bound_p": b_p, "bound_q": b_q, "bound_x": b_x, "bound_y": b_y,
        "bound_v": b_v, "bound_u": b_u, "bound_w": b_w, "bound_S": b_s,
        "soc_chain_max_kWh": float(chain),
        "soc_end_mismatch_kWh": float(abs(S[-1][-1] - run["S_end"])),
        "settle_flow_max_yuan": flow,
        "official_days": int(uniform_slice_count(n, offset)),
    }


def uniform_slice_count(n: int, offset: int) -> int:
    return max(0, n - offset)
