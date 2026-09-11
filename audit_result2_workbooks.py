"""Read-only structural and numerical audit for the official Q2 workbook."""
from __future__ import annotations

import argparse
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ERROR_LITERALS = {
    "#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A",
}
EXPECTED_SHEETS = ["计划购电量", "充放电量", "紧急购电量"]


def normalize(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def defined_name_signature(workbook) -> list[tuple[Any, ...]]:
    return sorted(
        (
            item.name,
            item.attr_text,
            item.localSheetId,
            item.hidden,
            item.function,
            item.vbProcedure,
            item.xlm,
            item.functionGroupId,
            item.shortcutKey,
            item.publishToServer,
            item.workbookParameter,
        )
        for item in workbook.defined_names.values()
    )


def validation_signature(worksheet) -> list[tuple[Any, ...]]:
    if worksheet.data_validations is None:
        return []
    return sorted(
        (
            str(item.sqref), item.type, item.operator, item.formula1, item.formula2,
            item.allow_blank, item.showDropDown, item.showInputMessage,
            item.showErrorMessage, item.errorStyle, item.error, item.errorTitle,
            item.prompt, item.promptTitle,
        )
        for item in worksheet.data_validations.dataValidation
    )


def workbook_snapshot(path: Path) -> tuple[Any, dict[str, Any]]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    error_cells: list[str] = []
    formula_cells: list[str] = []
    sheet_snapshot: dict[str, Any] = {}
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "e" or (
                    isinstance(cell.value, str) and cell.value.strip() in ERROR_LITERALS
                ):
                    error_cells.append(f"{worksheet.title}!{cell.coordinate}")
                if cell.data_type == "f":
                    formula_cells.append(f"{worksheet.title}!{cell.coordinate}={cell.value}")
        sheet_snapshot[worksheet.title] = {
            "dimension": worksheet.calculate_dimension(),
            "max_row": worksheet.max_row,
            "max_column": worksheet.max_column,
            "header": [normalize(cell.value) for cell in worksheet[1]],
            "merged_cells": sorted(str(item) for item in worksheet.merged_cells.ranges),
            "tables": sorted(worksheet.tables.keys()),
            "validations": validation_signature(worksheet),
            "freeze_panes": str(worksheet.freeze_panes) if worksheet.freeze_panes else None,
            "auto_filter": str(worksheet.auto_filter.ref) if worksheet.auto_filter.ref else None,
        }
    snapshot = {
        "sheet_names": workbook.sheetnames,
        "defined_names": defined_name_signature(workbook),
        "sheets": sheet_snapshot,
        "formula_cells": sorted(formula_cells),
        "error_cells": sorted(error_cells),
    }
    return workbook, snapshot


def compare_all_values(left, right) -> list[str]:
    differences: list[str] = []
    if left.sheetnames != right.sheetnames:
        return ["sheet order differs"]
    for sheet_name in left.sheetnames:
        first = left[sheet_name]
        second = right[sheet_name]
        if (first.max_row, first.max_column) != (second.max_row, second.max_column):
            differences.append(f"{sheet_name}: dimensions differ")
            continue
        for row in first.iter_rows():
            for cell in row:
                other = second[cell.coordinate]
                if normalize(cell.value) != normalize(other.value) or cell.data_type != other.data_type:
                    differences.append(f"{sheet_name}!{cell.coordinate}")
                    if len(differences) >= 20:
                        return differences
    return differences


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", default=".")
    parser.add_argument("--output", default="results/q2_strategy/result2终检审计.json")
    args = parser.parse_args()
    root = Path(args.project_root).resolve()
    template_path = root / "data" / "附件" / "附件5" / "result2.xlsx"
    official_path = root / "results" / "result2.xlsx"
    experiment_path = root / "results" / "q2_strategy" / "result2.xlsx"

    template, template_snapshot = workbook_snapshot(template_path)
    official, official_snapshot = workbook_snapshot(official_path)
    experiment, experiment_snapshot = workbook_snapshot(experiment_path)

    plan = official["计划购电量"]
    plan_rows = [row for row in plan.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    dates = [normalize(row[0])[:10] for row in plan_rows]
    contract_numeric = all(
        isinstance(value, (int, float)) and not isinstance(value, bool)
        for row in plan_rows for value in row[1:145]
    )
    structure_keys = ("sheet_names", "defined_names")
    workbook_structure_preserved = all(
        official_snapshot[key] == template_snapshot[key] for key in structure_keys
    ) and all(
        official_snapshot["sheets"][name][key] == template_snapshot["sheets"][name][key]
        for name in template_snapshot["sheet_names"]
        for key in ("header", "merged_cells", "tables", "validations", "freeze_panes", "auto_filter")
    )
    formula_structure_preserved = official_snapshot["formula_cells"] == template_snapshot["formula_cells"]
    result_differences = compare_all_values(official, experiment)
    checks = {
        "expected_sheet_names_and_order": official.sheetnames == EXPECTED_SHEETS,
        "template_workbook_structure_preserved": workbook_structure_preserved,
        "template_formula_structure_preserved": formula_structure_preserved,
        "no_excel_error_cells": not official_snapshot["error_cells"] and not experiment_snapshot["error_cells"],
        "official_and_experiment_workbooks_identical": not result_differences,
        "exactly_334_plan_days": len(plan_rows) == 334,
        "date_range_is_2025_02_01_to_2025_12_31": bool(dates) and dates[0] == "2025-02-01" and dates[-1] == "2025-12-31",
        "all_144_contract_slots_numeric": contract_numeric,
    }
    report = {
        "status": "PASS" if all(checks.values()) else "FAIL",
        "paths": {
            "template": str(template_path),
            "official": str(official_path),
            "experiment": str(experiment_path),
        },
        "sheet_names": official.sheetnames,
        "sheet_dimensions": {
            name: official_snapshot["sheets"][name]["dimension"] for name in official.sheetnames
        },
        "defined_name_count": len(official_snapshot["defined_names"]),
        "formula_cell_count": len(official_snapshot["formula_cells"]),
        "data_validation_count": sum(
            len(official_snapshot["sheets"][name]["validations"]) for name in official.sheetnames
        ),
        "error_cells": official_snapshot["error_cells"],
        "written_day_count": len(plan_rows),
        "first_date": dates[0] if dates else None,
        "last_date": dates[-1] if dates else None,
        "official_experiment_value_differences": result_differences,
        "checks": checks,
    }
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    output_path = (root / args.output).resolve()
    if output_path == root or root not in output_path.parents:
        raise ValueError("--output must resolve inside --project-root")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(rendered + "\n", encoding="utf-8")
    print(rendered)
    return 0 if report["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
