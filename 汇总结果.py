# -*- coding: utf-8 -*-
"""整理现有真实运行结果；不重新求解、不画图。"""
from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent
CSV = ROOT / "results" / "csv"
RES = ROOT / "results"


def read(name: str) -> pd.DataFrame:
    return pd.read_csv(CSV / name)


def clean_tables() -> dict[str, pd.DataFrame]:
    q2s = read("问题2_配置敏感性.csv").drop(columns=["daily_cash"], errors="ignore")
    q2r = read("问题2_参照方案.csv").drop(columns=["daily_cash"], errors="ignore")
    q3a = read("问题3_信息消融.csv")
    q3a.loc[q3a["tag"] == "abl_none", "获准集合S"] = "{}"
    q3c = read("问题3_方案对比.csv")
    q4e = read("问题4_价格基线误差.csv")
    q4c = read("问题4_方案对比.csv")
    st = read("模型评价_结构敏感性.csv")
    ps = read("模型评价_误差倍率压力测试.csv")
    common = ["group", "scenario", "eta_c", "eta_d", "capacity_kWh",
              "power_kW", "S0_kWh", "cash_total", "相对基准费用变化_元",
              "相对基准费用变化_pct", "contract_kwh", "emergency_kwh",
              "emergency_slots", "charge_kwh", "discharge_kwh",
              "unused_contract_kwh", "S_min", "S_max", "S_end", "numeric_ok"]
    st = st[common]
    ps = ps[["gamma"] + common]
    return {"Q2参数敏感性": q2s, "Q2参照方案": q2r,
            "Q3八组信息消融": q3a, "Q3方案对比": q3c,
            "Q4价格预测误差": q4e, "Q4方案与理想基准": q4c,
            "结构敏感性": st, "误差倍率压力测试": ps}


def key_summary(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    q1 = read("问题1_汇总.csv").iloc[0]
    q2 = tables["Q2参数敏感性"].query("预测器 == 'F1' and W == 30 and alpha == 0.9").iloc[0]
    q3 = tables["Q3八组信息消融"].query("tag == 'main'").iloc[0]
    q4 = tables["Q4方案与理想基准"].set_index("tag")
    return pd.DataFrame([
        {"问题": "Q1", "方案": "确定性LP", "现金费用_元": q1["全天购电费_元"],
         "合同购电量_kWh": q1["全天计划购电量_kWh"], "应急购电量_kWh": 0.0},
        {"问题": "Q2", "方案": "F1-W30-alpha0.9", "现金费用_元": q2["cash_total"],
         "合同购电量_kWh": q2["contract_kwh"], "应急购电量_kWh": q2["emergency_kwh"]},
        {"问题": "Q3", "方案": "S={6,12,18}", "现金费用_元": q3["现金总费_元"],
         "合同购电量_kWh": q3["最终合同量_kWh"], "应急购电量_kWh": q3["应急购电量_kWh"]},
        {"问题": "Q4-2", "方案": "median30固定合同", "现金费用_元": q4.loc["Q4-2_median30", "cash_total"],
         "合同购电量_kWh": q4.loc["Q4-2_median30", "contract_kwh"],
         "应急购电量_kWh": q4.loc["Q4-2_median30", "emergency_kwh"]},
        {"问题": "Q4-3", "方案": "median30可调合同", "现金费用_元": q4.loc["Q4-3_median30", "cash_total"],
         "合同购电量_kWh": q4.loc["Q4-3_median30", "contract_kwh"],
         "应急购电量_kWh": q4.loc["Q4-3_median30", "emergency_kwh"]},
    ])


def autosize(path: Path) -> None:
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            letter = col[0].column_letter
            width = min(28, max(10, max(len(str(c.value or "")) for c in col) + 2))
            ws.column_dimensions[letter].width = width
    wb.save(path)


def validate_answer_workbooks() -> dict:
    names = ["result1.xlsx", "result2.xlsx", "result3.xlsx",
             "result4-2.xlsx", "result4-3.xlsx"]
    details = []
    all_ok = True
    for name in names:
        template = openpyxl.load_workbook(ROOT / "data" / "附件" / "附件5" / name,
                                          data_only=False)
        result = openpyxl.load_workbook(RES / name, data_only=False)
        same_sheets = template.sheetnames == result.sheetnames
        same_dimensions = all(template[s].calculate_dimension() ==
                              result[s].calculate_dimension()
                              for s in template.sheetnames)
        errors, nonfinite, annual_complete = [], [], True
        for ws in result.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("#"):
                        errors.append(f"{ws.title}!{cell.coordinate}:{value}")
                    if isinstance(value, (int, float)) and not math.isfinite(float(value)):
                        nonfinite.append(f"{ws.title}!{cell.coordinate}")
            if ws.max_row == 335 and ws.max_column == 147:
                annual_complete = annual_complete and all(
                    ws.cell(r, c).value is not None
                    for r in range(2, 336) for c in range(1, 148))
        ok = same_sheets and same_dimensions and not errors and not nonfinite and annual_complete
        all_ok = all_ok and ok
        details.append({"file": name, "same_sheets": same_sheets,
                        "same_dimensions": same_dimensions,
                        "annual_cells_complete": annual_complete,
                        "excel_errors": errors, "nonfinite_cells": nonfinite,
                        "ok": ok})
    return {"all_ok": all_ok, "details": details}


def main() -> None:
    tables = clean_tables()
    summary = key_summary(tables)
    summary.to_csv(CSV / "全题关键指标.csv", index=False, encoding="utf-8-sig")
    tables["Q2参数敏感性"].to_csv(CSV / "问题2_配置敏感性.csv", index=False, encoding="utf-8-sig")
    tables["Q2参照方案"].to_csv(CSV / "问题2_参照方案.csv", index=False, encoding="utf-8-sig")
    tables["Q3八组信息消融"].to_csv(CSV / "问题3_信息消融.csv", index=False, encoding="utf-8-sig")
    tables["结构敏感性"].to_csv(CSV / "模型评价_结构敏感性.csv", index=False, encoding="utf-8-sig")
    tables["误差倍率压力测试"].to_csv(CSV / "模型评价_误差倍率压力测试.csv", index=False, encoding="utf-8-sig")
    out = RES / "补充分析结果.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="全题关键指标", index=False)
        for name, df in tables.items():
            df.to_excel(writer, sheet_name=name, index=False)
    autosize(out)
    answer_check = validate_answer_workbooks()
    report = {"workbook": str(out), "sheets": 1 + len(tables),
              "rows": {"全题关键指标": len(summary),
                       **{k: len(v) for k, v in tables.items()}},
              "finite_numeric_nonmissing": bool(all(
                  np.isfinite(a[~np.isnan(a)]).all()
                  for df in [summary, *tables.values()]
                  for a in [df.select_dtypes(include=[np.number]).to_numpy()])),
              "official_answer_files": ["result1.xlsx", "result2.xlsx", "result3.xlsx",
                                        "result4-2.xlsx", "result4-3.xlsx"],
              "official_answer_check": answer_check}
    with open(CSV / "结果文件校验.json", "w", encoding="utf-8") as fh:
        json.dump(report, fh, ensure_ascii=False, indent=2)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
